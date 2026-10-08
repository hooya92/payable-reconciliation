"""Synthetic tests for same-template, editable month-end statement drafts."""
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from adapters.excel.month_end_statement import write_month_end_statement
from adapters.excel.reader import read_prior
from domain.models import JournalLine, NewPayable, PayableItem, ReconcileResult, SourceRef, Status
from domain.period import AccountingPeriod


TARGET=AccountingPeriod(2026,8)


def payable(code, name, desc, amount, row):
    return PayableItem(code,name,desc,Decimal(amount),"2026-07-22",row,
                       SourceRef("prior.xlsx","26.07",row,"담당자"))


def write_flat(path):
    wb=Workbook()
    ws=wb.active
    ws.title="26.07"
    ws.merge_cells("A1:E1")
    ws["A1"]="2026년 7월 말 미지급금 명세서"
    ws.append(["거래처코드","거래처명","날짜","적요","금액"])
    ws.append(["000011","가상A","2026-07-22","이월 비용",100])
    ws.append(["000022","가상B","2026-07-22","지급완료",200])
    ws.append(["000033","가상C","2026-07-22","검토 대상",300])
    ws["E3"].fill=PatternFill(fill_type="solid",fgColor="DDEEFF")
    ws.column_dimensions["D"].width=44
    wb.save(path)


def write_grouped(path):
    wb=Workbook()
    ws=wb.active
    ws.title="26.07"
    ws.append(["2026-07-31 거래처 명세"])
    ws.append(["코드","거래처","날짜","적요","금액"])
    ws.append(["","","2026-07-22","운송비",100])
    ws.append(["","","2026-07-23","용역비",200])
    ws.append(["000011","가상A","소계","",300])
    wb.save(path)


class TemplateMonthEndTests(unittest.TestCase):
    def test_flat_template_has_clean_rows_and_removable_audit_sheets(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/"prior.xlsx"
            output=Path(folder)/"draft.xlsx"
            write_flat(path)
            before=path.read_bytes()
            p1=payable("000011","가상A","이월 비용","100",3)
            p2=payable("000022","가상B","지급완료","200",4)
            p3=payable("000033","가상C","검토 대상","300",5)
            new=NewPayable(
                JournalLine("000044","가상D","25301","미지급금","신규 비용",
                            Decimal(0),Decimal(400),"2026-08-05",9),
                Decimal(400),"신규 대변 확인",True
            )
            results=[
                ReconcileResult(Status.UNPAID,p1,None,"미지급","NO_DEBIT"),
                ReconcileResult(Status.MATCHED,p2,None,"대사완료","EXACT"),
                ReconcileResult(Status.INPUT_TYPO_SUSPECT,p3,None,"코드 오타 의심","CODE_TYPO"),
            ]
            write_month_end_statement(output,[path],results,[new],[],TARGET)
            self.assertEqual(path.read_bytes(),before)
            wb=load_workbook(output)
            try:
                self.assertEqual(wb.sheetnames,["26.08","대사내역","검토필요","변경내역"])
                ws=wb["26.08"]
                self.assertIn("2026년 8월",ws["A1"].value)
                self.assertIn("A1:E1",[str(x) for x in ws.merged_cells.ranges])
                self.assertEqual(ws.column_dimensions["D"].width,44)
                data=list(ws.iter_rows(min_row=3,values_only=True))
                self.assertEqual({r[3] for r in data},{"이월 비용","지급완료","검토 대상","신규 비용"})
                self.assertEqual(ws["F2"].value,"처리상태")
                self.assertEqual(sum(r[4] for r in data),1000)
                self.assertEqual(len(list(wb["검토필요"].values)),2)
                self.assertIn("당월 발생",[r[0] for r in list(wb["변경내역"].values)[1:]])
                self.assertEqual(ws.cell(3,5).fill.fgColor.rgb,"00DDEEFF")
            finally:
                wb.close()

    def test_grouped_template_subtotals_are_formula_based_and_reimportable(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/"prior.xlsx"
            o=Path(folder)/"draft.xlsx"
            write_grouped(p)
            prior=payable("000011","가상A","운송비","100",3)
            other=payable("000011","가상A","용역비","200",4)
            r=[
                ReconcileResult(Status.UNPAID,prior,None,"미지급","NO_DEBIT"),
                ReconcileResult(Status.UNPAID,other,None,"미지급","NO_DEBIT"),
            ]
            write_month_end_statement(o,[p],r,[],[],TARGET)
            wb=load_workbook(o)
            try:
                ws=wb["26.08"]
                self.assertEqual(ws["A1"].value,"2026-08-31 거래처 명세")
                self.assertEqual(ws["C5"].value,"소계")
                self.assertEqual(ws["A5"].value,"000011")
                self.assertEqual(ws["E5"].value,"=SUM(E3:E4)")
                self.assertFalse(any(r[0] for r in list(ws.values)[2:4]))
            finally:
                wb.close()
            read=read_prior(o,period=TARGET)
            self.assertEqual(len(read.items),2)
            self.assertEqual(len(read.issues),0)

    def test_real_group_subtotal_style_moves_with_insertions_and_deletions(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/"prior.xlsx"; o=Path(folder)/"draft.xlsx"
            wb=Workbook()
            ws=wb.active;ws.title="26.07"
            ws.append(["가상 물류 거래처별 명세"])
            ws.append(["코드","거래처","날짜","적요","금액"])
            ws.append(["","","2026-07-31","7월 지급완료",100])
            ws.append(["","","2026-07-31","7월 미지급",200])
            ws.append(["001111","가상A","소계","",300])
            ws.append(["","","2026-07-31","7월 모두지급",80])
            ws.append(["002222","가상B","소계","",80])
            yellow=PatternFill(fill_type="solid",fgColor="FFF2A6")
            for row in (5,7):
                for c in ws[row][:5]:
                    c.fill=yellow
                    c.font=c.font.copy(name="굴림",bold=True)
            ws["D3"].font=ws["D3"].font.copy(name="굴림",size=10)
            wb.save(p)
            a1=payable("001111","가상A","7월 지급완료","100",3)
            a2=payable("001111","가상A","7월 미지급","200",4)
            b1=payable("002222","가상B","7월 모두지급","80",6)
            fresh=NewPayable(
                JournalLine("003333","가상신규","25301","미지급금",
                            "8월 새 운임",Decimal(0),Decimal(50),"2026-08-21",9),
                Decimal(50),"당월 신규 발생",True,
            )
            write_month_end_statement(
                o,[p],[ReconcileResult(Status.MATCHED,a1),
                       ReconcileResult(Status.UNPAID,a2),
                       ReconcileResult(Status.MATCHED,b1)],
                [fresh],[],TARGET,
            )
            wb=load_workbook(o)
            try:
                ws=wb["26.08"]
                self.assertEqual(ws.max_row,9)
                self.assertEqual([ws.cell(r,3).value for r in (5,7,9)],["소계"]*3)
                self.assertEqual([ws.cell(r,1).value for r in (5,7,9)],["001111","002222","003333"])
                self.assertEqual([ws.cell(r,5).value for r in (3,4,6,8)],[100,200,80,50])
                self.assertEqual(ws["C8"].value,"2026-08-21")
                self.assertEqual(ws["E5"].value,"=SUM(E3:E4)")
                self.assertEqual(ws["E7"].value,"=SUM(E6:E6)")
                self.assertEqual(ws["E9"].value,"=SUM(E8:E8)")
                for row in (5,7,9):
                    self.assertEqual(ws.cell(row,5).fill.fgColor.rgb,"00FFF2A6")
                    self.assertEqual(ws.cell(row,5).font.name,"굴림")
                self.assertEqual(ws["D3"].font.name,"굴림")
                self.assertEqual(ws["D8"].font.name,"굴림")
            finally:
                wb.close()

    def test_uncertain_new_credit_is_excluded_and_listed(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/"prior.xlsx"; o=Path(folder)/"draft.xlsx"
            write_flat(p)
            new=NewPayable(
                JournalLine("000055","가상E","25301","미지급금","미확정",
                            Decimal(0),Decimal(500),"2026-08-06",9),
                Decimal(500),"당월 지급/조정 확인 필요",False
            )
            write_month_end_statement(o,[p],[],[new],[],TARGET)
            wb=load_workbook(o)
            try:
                self.assertEqual(wb["26.08"].max_row,6)
                self.assertEqual(wb["26.08"]["F6"].value,"확인 필요")
                self.assertIn("당월 신규 확인 필요",[r[0] for r in list(wb["검토필요"].values)[1:]])
            finally:
                wb.close()

    def test_unresolved_generated_statement_is_blocked_on_next_month_reconciliation(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/"prior.xlsx"
            output=Path(folder)/"draft.xlsx"
            write_flat(p)
            review=payable("000033","가상C","검토 대상","300",5)
            write_month_end_statement(
                output,[p],[ReconcileResult(Status.INPUT_TYPO_SUSPECT,review,None,
                                           "거래처코드 오타 의심","CODE_TYPO")],[],[],TARGET
            )
            with self.assertRaisesRegex(ValueError,"검토필요 항목"):
                read_prior(output,period=TARGET)
            # Final approval remains a human action. Deleting the review tab
            # after correcting the main sheet allows normal next-month input.
            wb=load_workbook(output)
            wb.remove(wb["검토필요"])
            wb.save(output)
            self.assertEqual(len(read_prior(output,period=TARGET).items),3)

    def test_original_sheet_layout_and_unrelated_sheets_are_preserved(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/"prior.xlsx"; o=Path(folder)/"draft.xlsx"
            write_flat(p)
            wb=load_workbook(p)
            ws=wb["26.07"]
            ws.sheet_view.zoomScale=77
            ws.freeze_panes="C3"
            ws.print_options.horizontalCentered=True
            ws.print_area="A1:E5"
            ws.row_dimensions[4].height=37
            ws["E4"].fill=PatternFill(fill_type="solid",fgColor="ABCDEF")
            note=wb.create_sheet("기존 안내")
            note["A1"]="전월 원본 참고 자료"
            note["A1"].fill=PatternFill(fill_type="solid",fgColor="AACCEE")
            wb.save(p)
            before=load_workbook(p)
            original_merge=[str(r) for r in before["26.07"].merged_cells.ranges]
            before.close()
            item=payable("000011","가상A","이월 비용","100",3)
            write_month_end_statement(o,[p],[ReconcileResult(Status.UNPAID,item)],[],[],TARGET)
            wb=load_workbook(o)
            try:
                self.assertEqual(wb.sheetnames,["26.08","기존 안내","대사내역","검토필요","변경내역"])
                ws=wb["26.08"]
                self.assertEqual(ws["E3"].fill.fgColor.rgb,"00DDEEFF")
                self.assertEqual(ws.max_row,5)
                self.assertEqual(ws.sheet_view.zoomScale,77)
                self.assertEqual(ws.freeze_panes,"C3")
                self.assertIn("$A$1:$E$5",str(ws.print_area))
                self.assertEqual([str(r) for r in ws.merged_cells.ranges],original_merge)
                self.assertEqual(wb["기존 안내"]["A1"].value,"전월 원본 참고 자료")
                self.assertEqual(wb["기존 안내"]["A1"].fill.fgColor.rgb,"00AACCEE")
            finally:
                wb.close()

    def test_merged_cells_in_body_are_rejected_without_affecting_source(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/"prior.xlsx"; o=Path(folder)/"draft.xlsx"
            write_flat(p)
            wb=load_workbook(p)
            wb["26.07"].merge_cells("B4:C4")
            wb.save(p)
            initial=p.read_bytes()
            with self.assertRaisesRegex(ValueError,"병합 셀"):
                write_month_end_statement(o,[p],[],[],[],TARGET)
            self.assertEqual(p.read_bytes(),initial)
            self.assertFalse(o.exists())

    def test_does_not_overwrite_source_or_guess_an_already_present_month(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/"prior.xlsx"; write_flat(p)
            prior=payable("000011","가상A","이월 비용","100",3)
            with self.assertRaisesRegex(ValueError,"원본"):
                write_month_end_statement(p,[p],[ReconcileResult(Status.UNPAID,prior)],[],[],TARGET)
            wb=load_workbook(p)
            wb.create_sheet("26.08")
            wb.save(p)
            with self.assertRaisesRegex(ValueError,"이미 있습니다"):
                write_month_end_statement(Path(folder)/"output.xlsx",[p],[],[],[],TARGET)


if __name__=="__main__":
    unittest.main()
