import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from adapters.excel.month_end_statement import write_month_end_statement
from application.service import run_reconciliation
from domain.models import Status
from domain.period import AccountingPeriod


class ScopedInputIssueTests(unittest.TestCase):
    def run_case(self,folder,unknown_identity=False,text_payment=False):
        prior=Path(folder)/"prior.xlsx";raw=Path(folder)/"raw.xlsx"
        wb=Workbook();ws=wb.active;ws.title="26.07"
        ws.append(["코드","거래처","날짜","적요","금액"])
        ws.append([None,None,"2026-07-31","시설관리비",170000])
        ws.append([None,None,"2026-07-31","지급","-       170000"])
        ws["E3"].number_format="@"
        ws.append(["002222","가상시설관","소계",None,
                   "=SUM(E2:E3)" if text_payment else 170000])
        ws.append([None,None,"2026-07-31","운송비",500])
        ws.append([None,None,"2026-07-31","지급",-100])
        ws.append(["001111","물류","소계",None,"=SUM(E5:E6)"])
        ws.append([None,None,"2026-07-31","청소비",100])
        ws.append(["003333","청소","소계",None,100])
        if unknown_identity:
            ws.append([None,"미상 거래처","2026-07-31","코드 누락",5])
        wb.save(prior);wb.close()
        wb=Workbook();ws=wb.active
        ws.append(["기표일자","계정코드","계정과목명","거래처코드","거래처명","적요","차변","대변"])
        for code,name in (("002222","가상시설관"),("001111","물류"),("003333","청소")):
            ws.append(["2026-08-10","25301","미지급금",code,name,"8월 신규",0,10])
        wb.save(raw);wb.close()
        return run_reconciliation([prior],[raw],{"25301"},AccountingPeriod(2026,8))

    def test_known_vendor_subtotal_error_does_not_block_other_vendors(self):
        with tempfile.TemporaryDirectory() as folder:
            run=self.run_case(folder)
            self.assertTrue(run.issues)
            self.assertEqual({issue.vendor_code for issue in run.issues},{"002222"})
            fresh={item.journal.vendor_code:item for item in run.new_items}
            self.assertFalse(fresh["002222"].auto_carry)
            self.assertIn("해당 거래처",fresh["002222"].reason)
            self.assertTrue(fresh["003333"].auto_carry)
            signed=next(r for r in run.results if r.prior.vendor_code=="001111")
            self.assertEqual(signed.status,Status.SIGNED_NET_AUTO)

    def test_unknown_vendor_identity_retains_global_hold(self):
        with tempfile.TemporaryDirectory() as folder:
            run=self.run_case(folder,unknown_identity=True)
            self.assertTrue(any(not issue.vendor_code for issue in run.issues))
            self.assertTrue(all(not item.auto_carry for item in run.new_items))
            signed=next(r for r in run.results if r.prior.vendor_code=="001111")
            self.assertEqual(signed.status,Status.SIGNED_OPENING_REVIEW)

    def test_text_negative_with_valid_subtotal_is_not_an_input_error(self):
        with tempfile.TemporaryDirectory() as folder:
            run=self.run_case(folder,text_payment=True)
            self.assertEqual(run.issues,[])
            self.assertTrue(all(item.auto_carry for item in run.new_items))
            self.assertTrue(all(r.status!=Status.SIGNED_OPENING_REVIEW for r in run.results))

    def test_generated_status_limits_input_hold_to_known_vendor(self):
        with tempfile.TemporaryDirectory() as folder:
            run=self.run_case(folder)
            prior=Path(folder)/"prior.xlsx"
            raw=Path(folder)/"raw.xlsx"
            original={path:path.read_bytes() for path in (prior,raw)}
            output=Path(folder)/"result.xlsx"
            write_month_end_statement(
                output,[prior],run.results,run.new_items,run.issues,run.period,
                standalone_debits=run.standalone_debits,
                journal_items=run.journal_items,journal_sources=run.journal_sources,
                source_paths=[prior,raw]
            )
            wb=load_workbook(output)
            try:
                new_rows=[row for row in wb["26.08"].iter_rows()
                          if row[3].value=="8월 신규"]
                self.assertEqual(len(new_rows),3)
                self.assertIn("입력",new_rows[0][5].value)
                self.assertEqual([row[5].value for row in new_rows[1:]],
                                 ["당월 신규","당월 신규"])
                self.assertEqual({row[1].value for row in wb["검토필요"].iter_rows(min_row=2)
                                  if row[0].value=="입력 데이터 확인"},{"002222"})
                self.assertTrue(all(path.read_bytes()==data for path,data in original.items()))
            finally:
                wb.close()
