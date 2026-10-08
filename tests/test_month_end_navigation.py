import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from adapters.excel.month_end_statement import write_month_end_statement
from application.service import run_reconciliation
from domain.period import AccountingPeriod
from tests.test_month_end_offset_audit import make_prior, make_raw


class MonthEndNavigationTests(unittest.TestCase):
    def test_links_follow_inserted_rows_and_preserve_original_row_numbers(self):
        with tempfile.TemporaryDirectory() as folder:
            prior=Path(folder)/"opening.xlsx"
            raw=Path(folder)/"journal.xlsx"
            output=Path(folder)/"result.xlsx"
            make_prior(prior);make_raw(raw)
            before={p:p.read_bytes() for p in (prior,raw)}
            run=run_reconciliation([prior],[raw],{"25301"},AccountingPeriod(2026,8))
            write_month_end_statement(
                output,[prior],run.results,run.new_items,run.issues,run.period,
                standalone_debits=run.standalone_debits,
                journal_items=run.journal_items,journal_sources=run.journal_sources,
                source_paths=[prior,raw]
            )
            wb=load_workbook(output)
            try:
                # First pair adds a payment before the second original vendor.
                # Original row 5 therefore appears at row 6 in the new sheet.
                for sheet in ("대사내역","변경내역","검토필요"):
                    for row in wb[sheet].iter_rows(min_row=2,max_col=8):
                        if row[7].hyperlink:
                            link=row[7].hyperlink
                            self.assertIsNone(link.target)
                            target=link.location.split("!")[1]
                            self.assertEqual(wb["26.08"][target].value,row[3].value)
                audit=wb["대사내역"]["H2"]
                self.assertEqual(audit.value,3)
                self.assertEqual(audit.hyperlink.location,"'26.08'!D3")
                second_original=next(row for row in wb["변경내역"].iter_rows(min_row=2)
                                     if row[6].value==prior.name and row[7].value==5)
                self.assertEqual(second_original[7].hyperlink.location,"'26.08'!D6")
                orphan=wb["검토필요"]["H2"]
                self.assertEqual(orphan.value,3)
                self.assertEqual(orphan.hyperlink.location,"'26.08'!D8")
                self.assertIn("원본 파일 기준",wb["대사내역"]["H1"].comment.text)
                self.assertTrue(all(p.read_bytes()==data for p,data in before.items()))
            finally:
                wb.close()
