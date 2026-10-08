import tempfile
import unittest
from collections import Counter
from pathlib import Path
from adapters.excel.reader import read_douzone, read_prior
from domain.models import Status
from domain.period import AccountingPeriod
from domain.reconciliation import new_payables, reconcile
from tests.synthetic_excel_factory import build

class SyntheticExcelE2ETests(unittest.TestCase):
    def test_thousands_of_rows_against_golden_answer(self):
        with tempfile.TemporaryDirectory() as tmp:
            expected=build(tmp,prior_count=2400,raw_noise=7200)
            root=Path(tmp); prior_items=[]; issues=[]
            for p in sorted(root.glob("담당자*_미지급.xlsx")):
                r=read_prior(p,p.stem); prior_items.extend(r.items); issues.extend(r.issues)
            dz=read_douzone(root/"더존_Raw_2026_1년.xlsx",{"25301"},AccountingPeriod(2026,7))
            issues.extend(dz.issues)
            result=reconcile(prior_items,dz.items); counts=Counter(x.status for x in result)
            self.assertEqual(len(issues),0)
            self.assertEqual(len(prior_items),expected["prior_count"])
            self.assertEqual(counts[Status.MATCHED],expected["MATCHED"])
            self.assertEqual(counts[Status.DESCRIPTION_MISMATCH],expected["DESCRIPTION_MISMATCH"])
            self.assertEqual(sum(x.rule=="CODE_AMOUNT_UNIQUE_WITH_NOTE" for x in result),expected["MATCHED_WITH_NOTE"])
            self.assertEqual(counts[Status.UNPAID],expected["UNPAID"])
            self.assertEqual(counts[Status.VENDOR_MISMATCH],expected["VENDOR_MISMATCH"])
            self.assertEqual(len(new_payables(dz.items,result)),expected["new_payables"])

    def test_split_raw_files_behave_like_one_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            expected=build(tmp,prior_count=1200,raw_noise=12000)
            root=Path(tmp); source=root/"더존_Raw_2026_1년.xlsx"
            from openpyxl import load_workbook, Workbook
            wb=load_workbook(source,read_only=True,data_only=True); ws=wb.active
            rows=ws.iter_rows(values_only=True); header=next(rows); chunks=[[],[],[]]
            for i,row in enumerate(rows): chunks[i%3].append(row)
            wb.close(); source.unlink()
            for n,chunk in enumerate(chunks):
                out=Workbook(); ow=out.active; ow.title="전표출력"; ow.append(header)
                for row in chunk: ow.append(row)
                out.save(root/f"더존_Raw_part{n+1}.xlsx")
            prior=[]
            for fp in sorted(root.glob("담당자*_미지급.xlsx")): prior.extend(read_prior(fp,fp.stem).items)
            journal=[]
            for fp in sorted(root.glob("더존_Raw_part*.xlsx")): journal.extend(read_douzone(fp,{"25301"},AccountingPeriod(2026,7)).items)
            result=reconcile(prior,journal); counts=Counter(x.status for x in result)
            self.assertEqual(counts[Status.MATCHED],expected["MATCHED"])
            self.assertEqual(sum(x.rule=="CODE_AMOUNT_UNIQUE_WITH_NOTE" for x in result),expected["MATCHED_WITH_NOTE"])
            self.assertEqual(len(new_payables(journal,result)),expected["new_payables"])

if __name__=="__main__":
    unittest.main()
