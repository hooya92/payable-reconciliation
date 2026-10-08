import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from application.service import run_reconciliation
from domain.models import Status
from domain.period import AccountingPeriod


def make_prior(path):
    wb=Workbook(); ws=wb.active
    ws.append(["거래처코드","거래처명","날짜","적요","금액"])
    ws.append(["001001","가상상사","2026-07-31","7월 유류비",100000])
    wb.save(path)


def make_raw(path, debit=100000, credit=0, desc="7월 유류비"):
    wb=Workbook(); ws=wb.active
    ws.append(["기표일자","계정코드","계정과목명","거래처코드","거래처명","적요","차변","대변"])
    ws.append(["2026-07-05","25301","미지급금-일반","001001","가상상사",desc,debit,credit])
    wb.save(path)


class ServiceSafetyTests(unittest.TestCase):
    def test_overlapping_raw_exports_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            prior=root/"prior.xlsx"; raw1=root/"raw_a.xlsx"; raw2=root/"raw_b.xlsx"
            make_prior(prior); make_raw(raw1); make_raw(raw2)
            with self.assertRaisesRegex(ValueError,"Raw 파일 간 동일 전표"):
                run_reconciliation([prior],[raw1,raw2],{"25301"},AccountingPeriod(2026,7))

    def test_orphan_current_debit_is_reported_and_cannot_silently_disappear(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            prior=root/"prior.xlsx"; raw=root/"raw.xlsx"
            make_prior(prior)
            wb=Workbook(); ws=wb.active
            ws.append(["기표일자","계정코드","계정과목명","거래처코드",
                       "거래처명","적요","차변","대변"])
            ws.append(["2026-07-05","25301","미지급금-일반",
                       "009999","전월에 없는 업체","미확인 차변",40000,0])
            wb.save(raw)
            result=run_reconciliation([prior],[raw],{"25301"},AccountingPeriod(2026,7))
            self.assertEqual(len(result.issues),1)
            self.assertIn("대응하지 않는 차변",result.issues[0].reason)
            self.assertEqual(result.results[0].status,Status.RAW_INPUT_INCOMPLETE)

    def test_blank_account_code_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            prior=root/"prior.xlsx"; raw=root/"raw.xlsx"
            make_prior(prior); make_raw(raw)
            with self.assertRaisesRegex(ValueError,"계정코드"):
                run_reconciliation([prior],[raw],set(),AccountingPeriod(2026,7))


if __name__=="__main__":
    unittest.main()
