"""Regression checks for previous-month opening balances and current-month close."""
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app import infer_period_from_inputs
from application.service import run_reconciliation
from adapters.excel.writer import write_result
from domain.models import JournalLine, Status
from domain.period import AccountingPeriod
from domain.reconciliation import new_payables


def raw_line(description, debit=0, credit=0, vendor="051330"):
    return JournalLine(
        vendor, "가상물류", "25301", "미지급금-일반",
        description, Decimal(str(debit)), Decimal(str(credit)), "2026-08-20", 2,
    )


class MonthEndTests(unittest.TestCase):
    def test_period_inference_requires_previous_statement_month(self):
        july=AccountingPeriod(2026,7)
        august=AccountingPeriod(2026,8)
        self.assertEqual(infer_period_from_inputs([[july]], [[august]]), (august,"statement+raw"))
        self.assertEqual(infer_period_from_inputs([[august]], [[august]]), (None,"conflict"))
        self.assertEqual(infer_period_from_inputs([[AccountingPeriod(2025,12)]], [[AccountingPeriod(2026,1)]]),
                         (AccountingPeriod(2026,1),"statement+raw"))

    def test_new_credit_is_net_of_same_month_payment(self):
        new=new_payables([raw_line("8월 운송",credit=50),raw_line("8월 운송",debit=20)],[])
        self.assertEqual(len(new),1)
        self.assertTrue(new[0].auto_carry)
        self.assertEqual(new[0].remaining,Decimal("30"))
        full=new_payables([raw_line("8월 운송",credit=50),raw_line("8월 운송",debit=50)],[])
        self.assertEqual(full[0].remaining,Decimal("0"))

    def test_duplicate_or_unmatched_new_credit_needs_review(self):
        credit=raw_line("8월 운송",credit=50)
        duplicate=new_payables([credit,raw_line("8월 운송",credit=40)],[])
        self.assertTrue(all(not x.auto_carry for x in duplicate))
        mismatch=new_payables([credit,raw_line("다른 적요",debit=20)],[])
        self.assertFalse(mismatch[0].auto_carry)

    def test_july_statement_august_journal_produces_august_close(self):
        with tempfile.TemporaryDirectory() as folder:
            prior=Path(folder)/"opening.xlsx"
            raw=Path(folder)/"journal.xlsx"
            output=Path(folder)/"result.xlsx"
            wb=Workbook()
            ws=wb.active
            ws.title="2026-07"
            ws.append(["거래처코드","거래처명","적요","금액"])
            ws.append(["051330","가상물류","7월 운송",100])
            wb.save(prior)

            wb=Workbook()
            ws=wb.active
            ws.title="전표출력"
            ws.append(["결의일","결의No","순번","기표일자","기표번호","구분",
                       "코드","계정과목명","코드","거래처명","적요","차변","대변"])
            ws.append(["2026-08-05",1,1,"2026-08-05",1,"일반",
                       "25301","미지급금-일반","051330","가상물류","7월 운송",60,0])
            ws.append(["2026-08-20",2,1,"2026-08-20",2,"일반",
                       "25301","미지급금-일반","051330","가상물류","8월 운송",0,30])
            ws.append(["2026-07-05",3,1,"2026-07-05",3,"일반",
                       "25301","미지급금-일반","051330","가상물류","과거 전표",500,0])
            wb.save(raw)

            period=AccountingPeriod(2026,8)
            result=run_reconciliation([prior],[raw],{"25301"},period)
            self.assertEqual(result.prior_count,1)
            self.assertEqual(result.results[0].status,Status.PARTIAL)
            self.assertTrue(result.new_items[0].auto_carry)
            self.assertEqual(result.new_items[0].remaining,Decimal("30"))
            write_result(output,result.results,result.new_items,result.issues,[prior,raw],period.label)
            wb=load_workbook(output,read_only=True,data_only=True)
            try:
                rows=list(wb["당월말명세서 초안"].values)
                detail=[row for row in rows if row[0] in ("부분지급잔액","당월신규")]
                self.assertEqual({row[0]:row[5] for row in detail},
                                 {"부분지급잔액":40,"당월신규":30})
                self.assertEqual([row[5] for row in rows if row[0]=="전체합계"],[70])
            finally:
                wb.close()

            # An explicitly titled 8월 opening sheet is not a valid 7월 close.
            wb=Workbook()
            ws=wb.active
            ws.title="2026-08"
            ws.append(["거래처코드","거래처명","적요","금액"])
            ws.append(["051330","가상물류","8월 운송",100])
            wb.save(prior)
            with self.assertRaisesRegex(ValueError,"전월 명세서"):
                run_reconciliation([prior],[raw],{"25301"},period)


if __name__=="__main__":
    unittest.main()
