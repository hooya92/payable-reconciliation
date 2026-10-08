"""Real Excel cell-style and date-type preservation when inserting month-end rows."""
import tempfile
import unittest
from datetime import datetime, date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from adapters.excel.month_end_statement import write_month_end_statement
from domain.models import JournalLine, NewPayable, PayableItem, ReconcileResult, SourceRef, Status
from domain.period import AccountingPeriod


class FullTemplateStyleTests(unittest.TestCase):
    def test_copies_font_borders_alignment_number_formats_and_date_type(self):
        with tempfile.TemporaryDirectory() as folder:
            prior=Path(folder)/"original.xlsx"
            output=Path(folder)/"generated.xlsx"
            wb=Workbook()
            ws=wb.active
            ws.title="26.07"
            ws["A1"]="2026년 7월 가상 미지급금"
            ws.append(["코드","거래처","날짜","적요","금액"])
            ws.append([None,None,datetime(2026,7,31),"7월 운임",100])
            ws.append(["001111","가상물류","소계",None,100])
            border=Border(left=Side(style="thin",color="111111"),
                          right=Side(style="medium",color="333333"),
                          bottom=Side(style="dotted",color="555555"))
            for c in ws[3][:5]:
                c.font=Font(name="굴림",size=10,color="223344")
                c.border=border
                c.alignment=Alignment(horizontal="center",vertical="center")
            ws["C3"].number_format="yyyy/mm/dd"
            ws["E3"].number_format="#,##0;[Red]-#,##0"
            ws.row_dimensions[3].height=29
            for c in ws[4][:5]:
                c.font=Font(name="굴림",size=10,bold=True)
                c.fill=PatternFill("solid",fgColor="FFF2A6")
                c.border=Border(top=Side(style="double",color="001122"))
            ws.row_dimensions[4].height=25
            wb.save(prior)
            wb.close()
            original_bytes=prior.read_bytes()
            payable=PayableItem("001111","가상물류","7월 운임",Decimal(100),"2026-07-31",3,
                                SourceRef(prior.name,"26.07",3,"담당자"))
            payment=JournalLine("001111","가상물류","25301","미지급금","7월 운임",
                                Decimal(100),Decimal(0),"2026-08-05",22)
            new=NewPayable(
                JournalLine("002222","가상신규","25301","미지급금","8월 신규",
                            Decimal(0),Decimal(250),"2026-08-23",33),
                Decimal(250),"당월 신규",True
            )
            write_month_end_statement(
                output,[prior],[ReconcileResult(Status.MATCHED,payable,payment,"일치","EXACT")],
                [new],[],AccountingPeriod(2026,8)
            )
            self.assertEqual(prior.read_bytes(),original_bytes)
            wb=load_workbook(output)
            try:
                ws=wb["26.08"]
                self.assertEqual(ws.max_row,7)
                self.assertEqual([ws.cell(r,5).value for r in (3,4,6)],[100,-100,250])
                self.assertEqual([r[0] for r in list(wb["대사내역"].values)[1:]],["대사 일치 후보"])
                for i in (3,4,6):
                    cell=ws.cell(i,5)
                    self.assertEqual(cell.font.name,"굴림")
                    self.assertEqual(cell.font.sz,10)
                    self.assertEqual(cell.border.left.style,"thin")
                    self.assertEqual(cell.border.right.style,"medium")
                    self.assertEqual(cell.border.bottom.style,"dotted")
                    self.assertEqual(cell.alignment.horizontal,"center")
                    self.assertEqual(cell.number_format,"#,##0;[Red]-#,##0")
                    self.assertEqual(ws.row_dimensions[i].height,29)
                    self.assertEqual(ws.cell(i,3).number_format,"yyyy/mm/dd")
                    self.assertIsInstance(ws.cell(i,3).value,(datetime,date))
                self.assertEqual(ws["C4"].value.date() if isinstance(ws["C4"].value,datetime)
                                 else ws["C4"].value,date(2026,8,5))
                for i in (5,7):
                    subtotal=ws.cell(i,5)
                    self.assertEqual(subtotal.fill.fgColor.rgb,"00FFF2A6")
                    self.assertEqual(subtotal.border.top.style,"double")
                    self.assertTrue(subtotal.font.bold)
                    self.assertEqual(ws.row_dimensions[i].height,25)
                self.assertEqual(ws["E5"].value,"=SUM(E3:E4)")
                self.assertEqual(ws["E7"].value,"=SUM(E6:E6)")
                self.assertEqual(ws["B5"].value,"가상물류")
                self.assertEqual(ws["B7"].value,"가상신규")
                self.assertIsNone(ws["F2"].value)  # Flags belong in removable review sheet
                self.assertTrue(wb.calculation.fullCalcOnLoad)
            finally:
                wb.close()


if __name__=="__main__":
    unittest.main()
