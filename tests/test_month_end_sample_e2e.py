"""Synthetic end-to-end example resembling the grouped yellow-subtotal statement."""
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from adapters.excel.month_end_statement import write_month_end_statement
from application.service import run_reconciliation
from domain.models import Status
from domain.period import AccountingPeriod


class ExampleMonthEndFlowTests(unittest.TestCase):
    def test_grouped_statement_realistic_add_remove_match_and_fonts(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            prior=folder/"statement.xlsx"
            raw=folder/"douzone.xlsx"
            output=folder/"generated.xlsx"

            wb=Workbook(); ws=wb.active; ws.title="26.07"
            ws["A1"]="(주)가상 미지급(거래처) 세부명세"
            ws["E1"]="2026-07-31"
            ws.append([])  # row 2 (blank)
            ws.append(["코드","거래처","날짜","적요","금액"])
            previous=[
                ("001111","가상물류서비스",[("7월 운송비",120000),("7월 보관료",80000)]),
                ("002222","가상시설관리",[("7월 시설관리비",70000)]),
                ("003333","가상기기렌탈",[("7월 복합기 임차료",95000)]),
                ("004444","가상주차관리",[("7월 주차비",55000)]),
                ("005555","가상청소",[("7월 청소비",45000)]),
            ]
            for code,name,items in previous:
                start=ws.max_row+1
                for desc,amount in items:
                    ws.append(["","","2026-07-31",desc,amount])
                    for cell in ws[ws.max_row][:5]:
                        cell.font=Font(name="굴림",size=10)
                ws.append([code,name,"소계","",f"=SUM(E{start}:E{ws.max_row})"])
                for cell in ws[ws.max_row][:5]:
                    cell.fill=PatternFill("solid",fgColor="FFF2A6")
                    cell.font=Font(name="굴림",size=10,bold=True)
            source_bytes=wb.save(prior)
            wb.close()
            original=prior.read_bytes()

            entries=[
                ("2026-08-05","001111","가상물류서비스","7월 운송비",120000,0),
                ("2026-08-07","003333","가상기기렌탈","7월 복합기 임차료",95000,0),
                ("2026-08-10","004444","가상주차관리","7월 주차비",30000,0),
                ("2026-08-18","001111","가상물류서비스","8월 운송비",0,65000),
                ("2026-08-20","006666","가상소모품상사","8월 신규 소모품",0,150000),
                ("2026-08-23","006666","가상소모품상사","8월 추가 소모품",0,45000),
                ("2026-08-28","007777","가상경비서비스","8월 경비 용역",0,60000),
            ]
            wb=Workbook();ws=wb.active;ws.title="전표출력"
            ws.append(["결의일","결의No","순번","기표일자","기표번호","구분",
                       "코드","계정과목명","코드","거래처명","적요","차변","대변"])
            for i,(d,code,name,desc,debit,credit) in enumerate(entries,1):
                ws.append([d,i,1,d,i,"일반","25301","미지급금-일반",code,name,desc,debit,credit])
            wb.save(raw);wb.close()

            period=AccountingPeriod(2026,8)
            run=run_reconciliation([prior],[raw],{"25301"},period)
            self.assertEqual(sum(x.status==Status.MATCHED for x in run.results),2)
            self.assertEqual(sum(x.status==Status.UNPAID for x in run.results),3)
            self.assertEqual(sum(x.status==Status.PARTIAL for x in run.results),1)
            self.assertEqual(run.issues,[])
            self.assertEqual(len(run.new_items),4)
            self.assertTrue(all(x.auto_carry for x in run.new_items))

            write_month_end_statement(
                output,[prior],run.results,run.new_items,run.issues,period
            )
            self.assertEqual(prior.read_bytes(),original)
            wb=load_workbook(output)
            ws=wb["26.08"]
            try:
                self.assertEqual(ws.max_row,22)  # 12 detail rows + 7 subtotals + 3 header rows
                details=[ws.cell(i,4).value for i in range(4,ws.max_row+1)
                         if ws.cell(i,3).value!="소계"]
                self.assertIn("7월 운송비",details)
                self.assertIn("7월 복합기 임차료",details)
                self.assertIn("7월 주차비",details)
                self.assertIn("7월 청소비",details)
                self.assertIn("8월 추가 소모품",details)
                subtotals=[(ws.cell(i,1).value,ws.cell(i,5).value,i)
                            for i in range(4,ws.max_row+1)
                            if ws.cell(i,3).value=="소계"]
                self.assertEqual([x[0] for x in subtotals],
                                 ["001111","002222","003333","004444","005555","006666","007777"])
                self.assertEqual(ws["C5"].value,"2026-08-05")
                self.assertEqual(ws["E5"].value,-120000)
                self.assertEqual(sum(ws.cell(i,5).value for i in range(4,23)
                                     if ws.cell(i,3).value!="소계"),570000)
                for _,formula,idx in subtotals:
                    self.assertTrue(formula.startswith("=SUM(E"))
                    self.assertEqual(ws.cell(idx,5).fill.fgColor.rgb,"00FFF2A6")
                    self.assertEqual(ws.cell(idx,5).font.name,"굴림")
                self.assertEqual(ws["D4"].font.name,"굴림")
                self.assertTrue(any("부분지급" in str(row[0])
                                    for row in list(wb["검토필요"].values)[1:]))
            finally:
                wb.close()


if __name__=="__main__":
    unittest.main()
