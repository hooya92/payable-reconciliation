import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from adapters.excel.reader import detect_statement_periods, read_douzone, read_prior
from domain.period import AccountingPeriod

class ExcelEdgeCaseTests(unittest.TestCase):
    def test_prior_header_offset_subtotal_and_leading_zero_numeric_code(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"prior.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["보고서"]); ws.append([""]); ws.append(["거래처코드","거래처명","날짜","적요","금액"])
            ws.append([51330,"가상물류","2026-07-31","7월 운송비",86000000]); ws["A4"].number_format="000000"
            ws.append(["","소계","","",86000000]); wb.save(p)
            r=read_prior(p,"담당자A")
            self.assertEqual(len(r.items),1); self.assertEqual(r.items[0].vendor_code,"051330")
            self.assertEqual(len(r.issues),0)

    def test_douzone_duplicate_code_headers_and_target_month_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"raw.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["전표출력"]); ws.append(["기표일자","코드","계정과목명","코드","거래처명","적요","차변","대변"])
            ws.append(["2026-07-31",25301,"미지급금-일반",51330,"가상물류","7월",100,0]); ws["D3"].number_format="000000"
            ws.append(["2026-08-01",25301,"미지급금-일반",51330,"가상물류","8월",100,0]); ws["D4"].number_format="000000"
            wb.save(p)
            r=read_douzone(p,{"25301"},AccountingPeriod(2026,8))
            self.assertEqual(len(r.items),1); self.assertEqual(r.items[0].vendor_code,"051330"); self.assertEqual(r.items[0].date,"2026-08-01")

    def test_suspicious_amount_is_excluded_and_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"prior.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["거래처코드","거래처명","날짜","적요","금액"]); ws.append(["001234","가상상사","2026-07-31","비용","8600만원"]); wb.save(p)
            r=read_prior(p)
            self.assertEqual(len(r.items),0); self.assertEqual(len(r.issues),1)

    def test_prior_missing_vendor_code_is_reported_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"prior.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["거래처코드","거래처명","날짜","적요","금액"])
            ws.append(["","가상상사","2026-07-31","유류비",100000]); wb.save(p)
            r=read_prior(p)
            self.assertEqual(len(r.items),0)
            self.assertTrue(any(x.field=="거래처코드" for x in r.issues))

    def test_prior_negative_amount_is_review_not_normal_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"prior.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["거래처코드","거래처명","날짜","적요","금액"])
            ws.append(["001234","가상상사","2026-07-31","수정건",-100000]); wb.save(p)
            r=read_prior(p)
            self.assertEqual(len(r.items),0)
            self.assertTrue(any("0 이하" in x.reason for x in r.issues))

    def test_douzone_requires_explicit_account_code(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"raw.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["기표일자","계정코드","계정과목명","거래처코드","거래처명","적요","차변","대변"])
            ws.append(["2026-08-01","25301","미지급금-일반","001234","가상상사","유류비",100000,0]); wb.save(p)
            with self.assertRaises(ValueError):
                read_douzone(p,set(),AccountingPeriod(2026,8))

    def test_douzone_missing_account_identifier_fails_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"raw.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["기표일자","거래처코드","거래처명","적요","차변","대변"])
            ws.append(["2026-08-01","001234","가상상사","유류비",100000,0]); wb.save(p)
            with self.assertRaises(ValueError):
                read_douzone(p,{"25301"},AccountingPeriod(2026,8))

    def test_douzone_negative_or_two_sided_row_is_review_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"raw.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["기표일자","계정코드","계정과목명","거래처코드","거래처명","적요","차변","대변"])
            ws.append(["2026-08-01","25301","미지급금-일반","001234","가상상사","역분개",-100000,0])
            ws.append(["2026-08-02","25301","미지급금-일반","001234","가상상사","이상행",100000,100000]); wb.save(p)
            r=read_douzone(p,{"25301"},AccountingPeriod(2026,8))
            self.assertEqual(len(r.items),0)
            self.assertEqual(len(r.issues),2)

    def test_prior_formula_amount_is_quarantined(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"prior_formula.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["거래처코드","거래처명","날짜","적요","금액"])
            ws.append(["001234","가상상사","2026-07-31","유류비","=50000+50000"]); wb.save(p)
            r=read_prior(p)
            self.assertEqual(len(r.items),0)
            self.assertTrue(any("수식" in x.reason for x in r.issues))

    def test_douzone_formula_amount_is_quarantined(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"raw_formula.xlsx"; wb=Workbook(); ws=wb.active
            ws.append(["기표일자","계정코드","계정과목명","거래처코드","거래처명","적요","차변","대변"])
            ws.append(["2026-08-01","25301","미지급금-일반","001234","가상상사","유류비","=50000+50000",0]); wb.save(p)
            r=read_douzone(p,{"25301"},AccountingPeriod(2026,8))
            self.assertEqual(len(r.items),0)
            self.assertTrue(any(x.field=="차변" and "수식" in x.reason for x in r.issues))

    def test_unrelated_cover_sheet_is_ignored_but_data_sheet_is_read(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"prior_multi.xlsx"; wb=Workbook()
            cover=wb.active; cover.title="표지"; cover.append(["명세서 안내"]); cover.append(["작성자","가상담당"])
            ws=wb.create_sheet("실제데이터")
            ws.append(["거래처코드","거래처명","날짜","적요","금액"])
            ws.append(["001234","가상상사","2026-07-31","유류비",100000]); wb.save(p)
            r=read_prior(p)
            self.assertEqual(len(r.items),1)
            self.assertEqual(r.recognized_sheets,["실제데이터"])

    def test_prior_partial_data_like_sheet_fails_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"prior_changed.xlsx"; wb=Workbook(); ws=wb.active; ws.title="명세"
            ws.append(["거래처코드","거래처명","적요","청구액"])
            ws.append(["001234","가상상사","유류비",100000]); wb.save(p)
            with self.assertRaisesRegex(ValueError,"데이터 표처럼 보이지만"):
                read_prior(p)

    def test_douzone_partial_data_like_sheet_fails_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"raw_changed.xlsx"; wb=Workbook(); ws=wb.active; ws.title="전표출력"
            ws.append(["기표일자","계정코드","거래처명","적요","차변","금액"])
            ws.append(["2026-08-01","25301","가상상사","유류비",100000,0]); wb.save(p)
            with self.assertRaisesRegex(ValueError,"전표 표처럼 보이지만"):
                read_douzone(p,{"25301"},AccountingPeriod(2026,8))

    def test_grouped_statement_details_inherit_vendor_from_following_subtotal(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"monthly.xlsx"; wb=Workbook(); ws=wb.active; ws.title="26.07"
            ws.append(["거래처코드","거래처명","날짜","적요","금액"])
            ws.append(["","","2026-07-31","7월 FMC사업부(주류) 용차료/유류비",64004897])
            ws.append(["","","2026-07-31","7월 FMC사업부(일반) 용차료/유류비",19325426])
            ws.append(["051330","(주)신일로지스시스템","소계","",83330323])
            wb.save(p)
            r=read_prior(p,"담당자",AccountingPeriod(2026,7))
            self.assertEqual(len(r.items),2)
            self.assertEqual([x.vendor_code for x in r.items],["051330","051330"])
            self.assertEqual(sum(x.amount for x in r.items),83330323)
            self.assertEqual(r.recognized_sheets,["26.07"])
            self.assertEqual(len(r.issues),0)

    def test_grouped_statement_subtotal_mismatch_fails_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"bad_group.xlsx"; wb=Workbook(); ws=wb.active; ws.title="26.07"
            ws.append(["거래처코드","거래처명","날짜","적요","금액"])
            ws.append(["","","2026-07-31","운송비 A",100])
            ws.append(["","","2026-07-31","운송비 B",200])
            ws.append(["001234","가상물류","소계","",999])
            wb.save(p)
            r=read_prior(p,"",AccountingPeriod(2026,7))
            self.assertEqual(len(r.items),0)
            self.assertTrue(any("상세 합계" in x.reason for x in r.issues))

    def test_detect_statement_periods_from_monthly_tabs(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"multi_month.xlsx"; wb=Workbook(); june=wb.active; june.title="26.06"
            june.append(["거래처코드","거래처명","날짜","적요","금액"])
            june.append(["001001","6월업체","2026-06-30","6월 비용",100])
            july=wb.create_sheet("2026-07")
            july.append(["거래처코드","거래처명","날짜","적요","금액"])
            july.append(["002002","7월업체","2026-07-31","7월 비용",200])
            cover=wb.create_sheet("안내")
            cover.append(["월별 명세서 안내"])
            wb.save(p)
            self.assertEqual(
                detect_statement_periods(p),
                [AccountingPeriod(2026,6),AccountingPeriod(2026,7)],
            )

    def test_monthly_workbook_reads_only_requested_statement_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"multi_month.xlsx"; wb=Workbook(); july=wb.active; july.title="26.07"
            july.append(["거래처코드","거래처명","날짜","적요","금액"])
            july.append(["001001","7월업체","2026-07-31","7월 비용",100])
            aug=wb.create_sheet("26.08")
            aug.append(["거래처코드","거래처명","날짜","적요","금액"])
            aug.append(["002002","8월업체","2026-08-31","8월 비용",200])
            wb.save(p)
            r=read_prior(p,"",AccountingPeriod(2026,7))
            self.assertEqual(len(r.items),1)
            self.assertEqual(r.items[0].vendor_code,"001001")
            self.assertEqual(r.recognized_sheets,["26.07"])

    def test_multiple_month_sheets_without_requested_month_are_blocked(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"multi_month.xlsx"; wb=Workbook(); july=wb.active; july.title="26.07"
            july.append(["거래처코드","거래처명","날짜","적요","금액"])
            july.append(["001001","7월업체","2026-07-31","7월 비용",100])
            aug=wb.create_sheet("26.08")
            aug.append(["거래처코드","거래처명","날짜","적요","금액"])
            aug.append(["002002","8월업체","2026-08-31","8월 비용",200])
            wb.save(p)
            with self.assertRaisesRegex(ValueError,"현재 대상 회계월은 2026년 10월이므로 2026년 9월 명세서가 필요합니다"):
                read_prior(p,"",AccountingPeriod(2026,9))

if __name__=="__main__": unittest.main()
