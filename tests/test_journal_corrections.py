import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from adapters.excel.month_end_statement import write_month_end_statement
from application.service import run_reconciliation
from domain.models import JournalLine, Status
from domain.period import AccountingPeriod
from domain.reconciliation import journal_correction_reviews


PERIOD=AccountingPeriod(2026,9)


class JournalCorrectionTests(unittest.TestCase):
    def make_files(self,folder,raw_rows):
        prior=Path(folder)/"prior.xlsx";raw=Path(folder)/"raw.xlsx"
        wb=Workbook();ws=wb.active;ws.title="26.08"
        ws.append(["미지급금 명세"])
        ws.append(["코드","거래처","날짜","적요","금액"])
        ws.append([None,None,"2026-08-13","벨트 교체",700000])
        ws.append([None,None,"2026-08-31","도어 구매",200000])
        ws.append(["001111","가상공영","소계",None,"=SUM(E3:E4)"])
        ws.append([None,None,"2026-08-31","정상 비용",100])
        ws.append(["000002","정상 업체","소계",None,100])
        ws["E5"].number_format='#,##0;[Red](#,##0);"-";@'
        ws["E5"].fill=PatternFill("solid",fgColor="FFFF00")
        wb.save(prior);wb.close()
        wb=Workbook();ws=wb.active
        ws.append(["기표일자","계정코드","계정과목명","거래처코드","거래처명","적요","차변","대변"])
        for code,name,desc,debit,credit,when in raw_rows:
            ws.append([when,"25301","미지급금",code,name,desc,debit,credit])
        wb.save(raw);wb.close()
        return prior,raw

    def export(self,folder,raw_rows):
        prior,raw=self.make_files(folder,raw_rows)
        before={path:path.read_bytes() for path in (prior,raw)}
        run=run_reconciliation([prior],[raw],{"25301"},PERIOD)
        output=Path(folder)/"result.xlsx"
        write_month_end_statement(output,[prior],run.results,run.new_items,run.issues,PERIOD,
            journal_items=run.journal_items,journal_sources=run.journal_sources,
            standalone_debits=run.standalone_debits,source_paths=[prior,raw])
        self.assertTrue(all(path.read_bytes()==data for path,data in before.items()))
        return run,output

    def test_credit_cancellation_reissue_and_payment_keep_exact_zero_balance(self):
        rows=[
            ("001111","가상공영","도어 구매_오발행건",0,-200000,"2026-09-01"),
            ("001111","가상공영","도어 구매",0,200000,"2026-09-01"),
            ("001111","가상공영","도어 구매",200000,0,"2026-09-07"),
            ("001111","가상공영","벨트 교체",700000,0,"2026-09-21"),
            ("000002","정상 업체","정상 비용",100,0,"2026-09-21"),
            ("000003","신규 업체","9월 신규 비용",0,50,"2026-09-21"),
        ]
        with tempfile.TemporaryDirectory() as folder:
            run,output=self.export(folder,rows)
            self.assertEqual(run.issues,[])
            self.assertEqual(len(run.journal_items),6)
            self.assertTrue(all(r.status==Status.CORRECTION_SUSPECT for r in run.results
                                if r.prior.vendor_code=="001111"))
            self.assertEqual(next(r for r in run.results if r.prior.vendor_code=="000002").status,
                             Status.MATCHED)
            self.assertTrue(next(n for n in run.new_items if n.journal.vendor_code=="000003").auto_carry)
            wb=load_workbook(output)
            try:
                ws=wb["26.09"]
                self.assertEqual([ws.cell(r,5).value for r in range(3,9)],
                                 [700000,200000,-200000,200000,-200000,-700000])
                self.assertEqual(sum(ws.cell(r,5).value for r in range(3,9)),0)
                self.assertEqual(ws["E9"].value,"=SUM(E3:E8)")
                self.assertEqual(ws["E9"].number_format,'#,##0;[Red](#,##0);"0";@')
                self.assertEqual(ws["E9"].fill.fgColor.rgb,"00FFFF00")
                self.assertEqual(ws["F9"].value,"소계 0")
                self.assertTrue(all(ws.cell(r,6).value.startswith("오발행·취소 의심") for r in range(3,9)))
                self.assertIn("전월 동일 금액 발생 후보 1건",ws["F5"].comment.text)
                self.assertIn("양수 동일 금액 후보 1건",ws["F5"].comment.text)
                self.assertTrue(all(row[1].value!="001111" for row in wb["대사내역"].iter_rows(min_row=2)))
                self.assertTrue(any(row[0].value==Status.CORRECTION_SUSPECT.value
                                    for row in wb["검토필요"].iter_rows(min_row=2)))
            finally:
                wb.close()

    def test_negative_debit_is_positive_adjustment_not_a_payment(self):
        with tempfile.TemporaryDirectory() as folder:
            run,output=self.export(folder,[("001111","가상공영","지급 취소",-50,0,"2026-09-01")])
            self.assertEqual(run.issues,[])
            wb=load_workbook(output)
            try:
                self.assertEqual(wb["26.09"]["E5"].value,50)
                self.assertTrue(wb["26.09"]["F5"].value.startswith(Status.CORRECTION_SUSPECT.value))
            finally:
                wb.close()

    def test_positive_keyword_is_reviewed_but_normal_payment_pair_is_not(self):
        lines=[JournalLine("1","가상","25301","미지급금",desc,Decimal(d),Decimal(c),"2026-09-01",r)
               for r,(desc,d,c) in enumerate((("청구",0,100),("지급",100,0),("수정 재발행",0,200)),2)]
        reviews=journal_correction_reviews(lines)
        self.assertEqual(set(reviews),{id(lines[2])})

    def test_negative_only_vendor_without_opening_is_preserved_for_review(self):
        with tempfile.TemporaryDirectory() as folder:
            run,output=self.export(folder,[("000099","조정 업체","수정 전표",-50,0,"2026-09-01")])
            self.assertEqual(run.issues,[])
            self.assertEqual(len(run.journal_items),1)
            self.assertFalse(any(r.prior.vendor_code=="000099" for r in run.results))
            wb=load_workbook(output)
            try:
                row=next(row for row in wb["26.09"].iter_rows() if row[3].value=="수정 전표")
                self.assertEqual(row[4].value,50)
                self.assertTrue(row[5].value.startswith(Status.CORRECTION_SUSPECT.value))
                self.assertTrue(any(row[1].value=="000099" and row[0].value==Status.CORRECTION_SUSPECT.value
                                    for row in wb["검토필요"].iter_rows(min_row=2)))
            finally:
                wb.close()

    def test_signed_opening_keeps_corrected_net_without_auto_approval(self):
        with tempfile.TemporaryDirectory() as folder:
            prior,raw=self.make_files(folder,[("001111","가상공영","취소",0,-10,"2026-09-01")])
            wb=load_workbook(prior);ws=wb.active
            ws["E3"]=100;ws["E4"]=-40
            wb.save(prior);wb.close()
            run=run_reconciliation([prior],[raw],{"25301"},PERIOD)
            result=next(r for r in run.results if r.prior.vendor_code=="001111")
            self.assertEqual(result.status,Status.CORRECTION_SUSPECT)
            self.assertEqual(result.closing_balance,50)
            self.assertEqual(run.issues,[])

    def test_correction_vendor_cannot_consume_another_vendors_same_amount_payment(self):
        with tempfile.TemporaryDirectory() as folder:
            prior,raw=self.make_files(folder,[
                ("001111","가상공영","오발행",0,-200000,"2026-09-01"),
                ("000099","다른 업체","도어 구매",200000,0,"2026-09-02"),
            ])
            run=run_reconciliation([prior],[raw],{"25301"},PERIOD)
            self.assertTrue(all(r.journal is None for r in run.results
                                if r.prior.vendor_code=="001111"))
            self.assertTrue(any(line.vendor_code=="000099" for line in run.standalone_debits))

    def test_same_amount_candidates_are_scoped_to_vendor_account_and_same_side(self):
        def line(code,account,debit,credit,row):
            return JournalLine(code,"가상",account,"미지급금","청구",Decimal(debit),Decimal(credit),
                               "2026-09-01",row)
        cancellation=line("1","25301",0,-100,2)
        others=[line("2","25301",0,100,3),line("1","25302",0,100,4),
                line("1","25301",100,0,5)]
        self.assertNotIn("동일 금액 후보",journal_correction_reviews([cancellation]+others)[id(cancellation)])
        matches=[line("1","25301",0,100,6),line("1","25301",0,100,7)]
        reason=journal_correction_reviews([cancellation]+others+matches)[id(cancellation)]
        self.assertIn("동일 금액 후보 2건",reason)
        self.assertIn("연결 미확정",reason)

    def test_zero_subtotal_format_keeps_accounting_escapes_and_quoted_semicolons(self):
        from adapters.excel.month_end_status import write_status
        wb=Workbook();ws=wb.active
        source=r'_-* #,##0_-;\-* #,##0_-;_-* "-"??_-;"참고;내용"'
        ws["E2"].number_format=source;ws["E2"]="=SUM(E3:E4)"
        write_status(ws,2,6,5,"소계",subtotal=True,subtotal_zero=True)
        self.assertEqual(ws["E2"].number_format,
                         r'_-* #,##0_-;\-* #,##0_-;"0";"참고;내용"')
        self.assertEqual(ws["F2"].value,"소계 0")
        wb.close()
