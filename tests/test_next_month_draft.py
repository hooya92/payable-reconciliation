import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from adapters.excel.writer import write_result
from domain.models import JournalLine, NewPayable, PayableItem, ReconcileResult, SourceRef, Status

class NextMonthDraftTests(unittest.TestCase):
    def test_draft_keeps_uncertain_new_credits_for_review(self):
        src=SourceRef("담당자A.xlsx","명세",7,"담당자A")
        unpaid=PayableItem("001001","가상A","7월 운송비",Decimal("1000"),"2026-07-31",7,src)
        matched=PayableItem("001002","가상B","7월 보관비",Decimal("2000"),"2026-07-31",8,src)
        review=PayableItem("001003","가상C","7월 용역비",Decimal("3000"),"2026-07-31",9,src)
        j=JournalLine("001002","가상B","25301","미지급금-일반","7월 보관비",Decimal("2000"),Decimal("0"),"2026-07-10",20)
        new=NewPayable(JournalLine("001004","가상D","25301","미지급금-일반","8월 신규비",Decimal("0"),Decimal("4000"),"2026-08-31",30),Decimal("4000"),"입력 확인 필요",False)
        results=[
            ReconcileResult(Status.UNPAID,unpaid,None,"대응 차변 없음","NO_DEBIT"),
            ReconcileResult(Status.MATCHED,matched,j,"일치","EXACT"),
            ReconcileResult(Status.DESCRIPTION_MISMATCH,review,j,"적요 불일치","VENDOR_AMOUNT"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/"result.xlsx"
            write_result(out,results,[new],[],[], "2026년 8월")
            wb=load_workbook(out,data_only=True); ws=wb["당월말명세서 초안"]
            rows=list(ws.iter_rows(min_row=2,values_only=True))
            detail=[r for r in rows if r[0] in ("미지급이월","당월신규")]
            subtotals=[r for r in rows if r[0]=="소계"]
            totals=[r for r in rows if r[0]=="전체합계"]
            self.assertEqual(len(detail),1)
            self.assertEqual(detail[0][0],"미지급이월")
            self.assertEqual(detail[0][5],1000)
            self.assertNotIn(3000,{r[5] for r in detail})
            self.assertEqual(len(subtotals),1)
            self.assertEqual(len(totals),1)
            self.assertEqual(totals[0][5],1000)
            review=list(wb["당월신규명세"].iter_rows(min_row=2,values_only=True))
            self.assertEqual(len(review),1)
            self.assertEqual(review[0][0],"사람 확인 필요")
            self.assertEqual(review[0][6],4000)
            wb.close()

if __name__=="__main__": unittest.main()
