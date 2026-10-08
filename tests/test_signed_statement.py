"""Signed manager statements: positive accruals, negative adjustments and vendor subtotals."""
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from adapters.excel.reader import read_prior
from adapters.excel.writer import write_result
from application.service import run_reconciliation
from domain.models import Status
from domain.period import AccountingPeriod


def statement(path, subtotal=120, formula=False, include_clean=True):
    wb=Workbook()
    ws=wb.active
    ws.title="26.09"
    ws.append(["(주)가상 미지급(거래처) 세부명세",None,None,None,"2026-09-30"])
    ws.append(["코드","거래처","날짜","적요","금액"])
    ws.append(["","","2026-09-30","9월 운송비",100])
    ws.append(["","","2026-09-30","9월 보관비",60])
    ws.append(["","","2026-09-15","8월분 지급",-40])
    ws.append(["001234","가상물류","소계","", "=SUM(E3:E5)" if formula else subtotal])
    if include_clean:
        ws.append(["","","2026-09-30","9월 통신비",70])
        ws.append(["009999","가상통신","소계","",70])
    wb.save(path)


class SignedStatementTests(unittest.TestCase):
    def test_signed_group_net_is_verified_but_not_assigned_to_invoices(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            statement(prior)
            result=read_prior(prior,"담당자",AccountingPeriod(2026,9))
            self.assertEqual(len(result.issues),0)
            self.assertEqual(len(result.review_items),1)
            balance,reason=result.review_items[0]
            self.assertEqual((balance.vendor_code,balance.amount),("001234",Decimal("120")))
            self.assertEqual(balance.source.row,6)
            self.assertIn("양수 160원 + 음수 -40원 = 소계 120원",reason)
            self.assertEqual(len(result.items),1)
            self.assertEqual((result.items[0].vendor_code,result.items[0].amount),
                             ("009999",Decimal("70")))

    def test_subtotal_sum_formula_is_independently_recomputed(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            statement(prior,formula=True)
            result=read_prior(prior,"담당자",AccountingPeriod(2026,9))
            self.assertEqual(len(result.issues),0)
            self.assertEqual(result.review_items[0][0].amount,Decimal("120"))
            self.assertEqual(len(result.items),1)

    def test_mismatched_signed_subtotal_never_carries_balance(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            statement(prior,subtotal=121,include_clean=False)
            result=read_prior(prior,"",AccountingPeriod(2026,9))
            self.assertEqual(result.items,[])
            self.assertEqual(result.review_items,[])
            self.assertTrue(any("상세 합계 120원과 소계 121원" in x.reason for x in result.issues))

    def test_zero_signed_subtotal_is_valid_no_outstanding_balance(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            statement(prior,include_clean=False,subtotal=0)
            wb=load_workbook(prior)
            wb.active["E5"]=-160
            wb.save(prior)
            result=read_prior(prior,"",AccountingPeriod(2026,9))
            self.assertEqual(result.items,[])
            self.assertEqual(len(result.review_items),1)
            self.assertEqual(result.review_items[0][0].amount,Decimal(0))
            self.assertEqual(result.issues,[])

    def test_unverified_subtotal_formula_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            statement(prior,include_clean=False)
            wb=load_workbook(prior)
            wb.active["E6"]="=SUM(E3:E4)"
            wb.save(prior)
            result=read_prior(prior,"",AccountingPeriod(2026,9))
            self.assertEqual(result.items,[])
            self.assertEqual(result.review_items,[])
            self.assertTrue(any(x.field=="수식" for x in result.issues))

    def test_flat_negative_blocks_same_vendor_positive_entries(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            wb=Workbook()
            ws=wb.active
            ws.title="26.09"
            ws.append(["코드","거래처","날짜","적요","금액"])
            ws.append(["001234","가상물류","2026-09-30","청구",100])
            ws.append(["001234","가상물류","2026-09-30","차감",-20])
            ws.append(["009999","가상통신","2026-09-30","통신",30])
            wb.save(prior)
            result=read_prior(prior,"",AccountingPeriod(2026,9))
            self.assertEqual([x.vendor_code for x in result.items],["009999"])
            self.assertTrue(any("순잔액 확인 전" in x.reason for x in result.issues))

    def test_signed_net_zero_is_auto_closed_and_negative_is_reviewed(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            raw=Path(directory)/"journal.xlsx"
            statement(prior,include_clean=False)
            wb=Workbook(); ws=wb.active
            ws.append(["기표일자","계정코드","계정과목명","거래처코드",
                       "거래처명","적요","차변","대변"])
            ws.append(["2026-10-02","25301","미지급금-일반",
                       "001234","가상물류","지급",120,0])
            wb.save(raw)
            zero=run_reconciliation([prior],[raw],{"25301"},AccountingPeriod(2026,10))
            self.assertEqual(zero.results[0].status,Status.SIGNED_NET_AUTO)
            self.assertEqual(zero.results[0].closing_balance,Decimal(0))
            output=Path(directory)/"zero.xlsx"
            write_result(output,zero.results,zero.new_items,zero.issues,[prior,raw],"2026년 10월")
            result=load_workbook(output,read_only=True,data_only=True)
            try:
                self.assertEqual(len(list(result["당월말명세서 초안"].values)),1)
            finally:
                result.close()
            wb=load_workbook(raw)
            wb.active.cell(2,7).value=130
            wb.save(raw)
            negative=run_reconciliation([prior],[raw],{"25301"},AccountingPeriod(2026,10))
            self.assertEqual(negative.results[0].status,Status.SIGNED_OPENING_REVIEW)
            self.assertEqual(negative.results[0].closing_balance,Decimal(-10))
            self.assertIn("음수 잔액",negative.results[0].reason)

    def test_zero_opening_subtotal_detects_october_overpayment(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            raw=Path(directory)/"journal.xlsx"
            statement(prior,include_clean=False,subtotal=0)
            wb=load_workbook(prior)
            wb.active["E5"]=-160
            wb.save(prior)

            wb=Workbook(); ws=wb.active
            ws.append(["기표일자","계정코드","계정과목명",
                       "거래처코드","거래처명","적요","차변","대변"])
            ws.append(["2026-10-03","25301","미지급금-일반",
                       "001234","가상물류","추가 지급",25,0])
            wb.save(raw)
            run=run_reconciliation([prior],[raw],{"25301"},AccountingPeriod(2026,10))
            self.assertEqual(run.prior_count,1)
            self.assertEqual(len(run.results),1)
            self.assertEqual(run.results[0].closing_balance,Decimal(-25))
            self.assertEqual(run.results[0].status,Status.SIGNED_OPENING_REVIEW)

    def test_invalid_raw_blocks_signed_net_auto_roll_forward(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            raw=Path(directory)/"journal.xlsx"
            statement(prior,include_clean=False)
            wb=Workbook(); ws=wb.active
            ws.append(["기표일자","계정코드","계정과목명","거래처코드",
                       "거래처명","적요","차변","대변"])
            ws.append(["2026-10-02","25301","미지급금-일반",
                       "001234","가상물류","10월 지급",10,0])
            ws.append(["2026-10-03","25301","미지급금-일반",
                       "001234","가상물류","수정분개",-5,0])
            wb.save(raw)
            run=run_reconciliation([prior],[raw],{"25301"},AccountingPeriod(2026,10))
            self.assertTrue(run.issues)
            self.assertEqual(run.results[0].status,Status.SIGNED_OPENING_REVIEW)
            self.assertIn("자동 반영 보류",run.results[0].reason)

    def test_signed_balance_is_carried_once_as_vendor_net(self):
        with tempfile.TemporaryDirectory() as directory:
            prior=Path(directory)/"prior.xlsx"
            raw=Path(directory)/"journal.xlsx"
            output=Path(directory)/"out.xlsx"
            statement(prior)
            wb=Workbook()
            ws=wb.active
            ws.title="전표출력"
            ws.append(["기표일자","계정코드","계정과목명","거래처코드",
                       "거래처명","적요","차변","대변"])
            ws.append(["2026-10-03","25301","미지급금-일반","009999",
                       "가상통신","10월 통신비",0,30])
            ws.append(["2026-10-07","25301","미지급금-일반","001234",
                       "가상물류","9월 운송비",100,0])
            ws.append(["2026-10-15","25301","미지급금-일반","001234",
                       "가상물류","10월 물류비",0,90])
            ws.append(["2026-09-01","25301","미지급금-일반","001234",
                       "가상물류","지난달 건",999,0])
            wb.save(raw)

            run=run_reconciliation([prior],[raw],{"25301"},AccountingPeriod(2026,10))
            self.assertEqual(run.prior_count,2)
            balance=[r for r in run.results if r.status==Status.SIGNED_NET_AUTO]
            self.assertEqual(len(balance),1)
            self.assertEqual(balance[0].prior.amount,Decimal("120"))
            self.assertEqual(balance[0].closing_balance,Decimal("110"))  # 120 + 90 - 100
            self.assertEqual(balance[0].prior.vendor_code,"001234")
            self.assertFalse(any(item.journal.vendor_code=="001234" for item in run.new_items))
            write_result(output,run.results,run.new_items,run.issues,[prior,raw],"2026년 10월")
            result=load_workbook(output,data_only=True,read_only=True)
            try:
                reviews=list(result["확인필요"].values)
                self.assertFalse(any(row[5]=="001234" for row in reviews[1:]))
                draft=list(result["당월말명세서 초안"].values)
                signed=[row for row in draft[1:] if row[1]=="001234" and row[0]!="소계"]
                self.assertEqual(len(signed),1)
                self.assertEqual(signed[0][0],"거래처순잔액")
                self.assertEqual(signed[0][5],110)
                self.assertIn("소계 120원 검증 완료",signed[0][6])
                self.assertEqual(len([row for row in draft[1:] if row[1]=="001234" and row[0]=="소계"]),1)
                self.assertTrue(any(row[1]=="009999" for row in draft[1:]))
                new_rows=list(result["당월신규명세"].values)
                self.assertFalse(any(row[3]=="001234" for row in new_rows[1:]))
                complete=list(result["자동대사완료"].values)
                self.assertTrue(any(row[1]=="001234" and row[4]==110 for row in complete[1:]))
            finally:
                result.close()


if __name__=="__main__":
    unittest.main()
