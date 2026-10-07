"""Small clean happy-path dataset for manual UI verification.

Run locally:
    python -m tests.synthetic_clean_factory

Generated Excel files are ignored by git and contain only synthetic data.
The statement intentionally uses the real-like monthly-tab + grouped-subtotal layout.
"""
from pathlib import Path
from openpyxl import Workbook


def build(out_dir="synthetic_clean"):
    root=Path(out_dir); root.mkdir(parents=True,exist_ok=True)

    prior=root/"정상_월별명세서.xlsx"
    wb=Workbook()

    # Keep another month in the same workbook to verify month-tab selection.
    june=wb.active; june.title="26.06"
    june.append(["거래처코드","거래처명","날짜","적요","금액"])
    june.append(["009999","가상과거업체","2026-06-30","6월 과거비용",123456])

    july=wb.create_sheet("26.07")
    july.append(["거래처코드","거래처명","날짜","적요","금액"])
    # Real-like grouped layout: detail rows first, vendor identity on following subtotal row.
    july.append(["","","2026-07-31","7월 FMC사업부(주류) 용차료/유류비",64004897])
    july.append(["","","2026-07-31","7월 FMC사업부(일반) 용차료/유류비",19325426])
    july.append(["051330","(주)가상로지스시스템","소계","",83330323])

    july.append(["","","2026-07-31","7월 클라우드 서비스 이용료",547350])
    july.append(["055866","주식회사 가상인포","소계","",547350])

    july.append(["","","2026-07-31","7월 보안용역료",25443033])
    july.append(["070070","(주)가상시큐리티","소계","",25443033])
    wb.save(prior)

    raw=root/"정상_더존Raw.xlsx"
    wb=Workbook(); ws=wb.active; ws.title="전표출력"
    ws.append(["기표일자","계정코드","계정과목명","거래처코드","거래처명","적요","차변","대변"])
    ws.append(["2026-07-03","25301","미지급금-일반","051330","(주)가상로지스시스템","7월 FMC사업부(주류) 용차료/유류비",64004897,0])
    ws.append(["2026-07-03","25301","미지급금-일반","051330","(주)가상로지스시스템","7월 FMC사업부(일반) 용차료/유류비",19325426,0])
    ws.append(["2026-07-05","25301","미지급금-일반","055866","주식회사 가상인포","7월 클라우드 서비스 이용료",547350,0])
    ws.append(["2026-07-07","25301","미지급금-일반","070070","(주)가상시큐리티","7월 보안용역료",25443033,0])

    # Noise outside target month/account must be ignored.
    ws.append(["2026-08-31","25301","미지급금-일반","099999","기간외가상","8월 기간외 전표",999999,0])
    ws.append(["2026-07-10","99999","기타계정","088888","타계정가상","타계정 전표",777777,0])
    wb.save(raw)

    return {
        "prior_path":prior,
        "raw_path":raw,
        "period_year":2026,
        "period_month":7,
        "expected_auto_match":4,
        "expected_review":0,
        "expected_input_issues":0,
        "expected_new_items":0,
    }


if __name__=="__main__":
    result=build()
    print(f"생성 완료: {Path(result['prior_path']).parent.resolve()}")
    print("대상 회계월: 2026년 7월")
    print("예상 결과: 자동대사 4 / 검토 필요 0 / 입력 확인 0 / 당월 신규 0")
