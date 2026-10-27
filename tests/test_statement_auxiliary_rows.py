import tempfile
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from adapters.excel.month_end_statement import write_month_end_statement
from adapters.excel.reader import read_prior
from domain.models import JournalLine, NewPayable, ReconcileResult, Status
from domain.period import AccountingPeriod
from tests.test_month_end_template import write_grouped, write_flat


class StatementAuxiliaryRowsTests(unittest.TestCase):
    def test_merged_total_reference_and_unknown_note_survive_new_raw_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            prior=Path(folder)/"prior.xlsx"; output=Path(folder)/"result.xlsx"
            write_grouped(prior)
            wb=load_workbook(prior);ws=wb.active
            ws.merge_cells("A6:D6");ws["A6"]="합 계";ws["E6"]="=SUM(E3:E5)/2"
            ws["E6"].fill=PatternFill("solid",fgColor="92D050")
            ws.row_dimensions[6].height=29
            ws["D7"]="더존";ws["E7"]=300
            ws["D8"]="추가로 작성한 업무 안내"
            ws.print_area="A1:E8"
            wb.save(prior);wb.close()
            before=prior.read_bytes()
            read=read_prior(prior,period=AccountingPeriod(2026,7))
            self.assertEqual(read.issues,[])
            self.assertEqual([item.amount for item in read.items],[Decimal(100),Decimal(200)])
            results=[ReconcileResult(Status.UNPAID,item) for item in read.items]
            line=JournalLine("000011","가상A","25301","미지급금","8월 신규",
                             Decimal(0),Decimal(50),"2026-08-10",2)
            fresh=NewPayable(line,Decimal(50),"",True)
            write_month_end_statement(output,[prior],results,[fresh],[],AccountingPeriod(2026,8),
                                      journal_items=[line])
            wb=load_workbook(output);ws=wb["26.08"]
            try:
                self.assertEqual(ws["E6"].value,"=SUM(E3:E5)")
                self.assertEqual(ws["A7"].value,"합 계")
                self.assertEqual(ws["E7"].value,'=SUMIF(C3:C6,"소계",E3:E6)')
                self.assertIn("A7:D7",{str(r) for r in ws.merged_cells.ranges})
                self.assertEqual(ws["E7"].fill.fgColor.rgb,"0092D050")
                self.assertEqual(ws.row_dimensions[7].height,29)
                self.assertEqual(ws["D8"].value,"더존")
                self.assertEqual(ws["E8"].value,300)
                self.assertIn("전월 원본",ws["E8"].comment.text)
                self.assertEqual(ws["D9"].value,"추가로 작성한 업무 안내")
                self.assertIn("$E$9",str(ws.print_area))
                preserved=[r for r in wb["변경내역"].iter_rows(min_row=2)
                           if r[0].value=="보조 행 보존"]
                self.assertEqual(len(preserved),3)
                self.assertEqual(preserved[0][7].hyperlink.location,"'26.08'!D7")
            finally:
                wb.close()
            self.assertEqual(prior.read_bytes(),before)
            reread=read_prior(output,period=AccountingPeriod(2026,8))
            self.assertEqual(reread.issues,[])
            self.assertEqual(sum(item.amount for item in reread.items),Decimal(350))

    def test_unknown_reference_formula_is_not_rebased_or_recalculated(self):
        with tempfile.TemporaryDirectory() as folder:
            prior=Path(folder)/"prior.xlsx";output=Path(folder)/"result.xlsx"
            write_flat(prior)
            wb=load_workbook(prior);ws=wb.active
            ws["D6"]="전체합계";ws["E6"]=600
            ws["D7"]="검증";ws["E7"]="=E6-123"
            wb.save(prior);wb.close()
            read=read_prior(prior,period=AccountingPeriod(2026,7))
            self.assertEqual(read.issues,[])
            write_month_end_statement(output,[prior],
                [ReconcileResult(Status.UNPAID,item) for item in read.items],[],[],
                AccountingPeriod(2026,8))
            wb=load_workbook(output)
            try:
                ws=wb["26.08"]
                self.assertEqual(ws["E6"].value,"=SUM(E3:E5)")
                self.assertEqual(ws["E7"].value,"=E6-123")
                self.assertEqual(ws["E7"].data_type,"s")
                self.assertIn("전월 원본 참고 수식",ws["E7"].comment.text)
            finally:
                wb.close()

    def test_note_inside_group_is_preserved_without_becoming_a_payable(self):
        with tempfile.TemporaryDirectory() as folder:
            prior=Path(folder)/"prior.xlsx";output=Path(folder)/"result.xlsx"
            write_grouped(prior)
            wb=load_workbook(prior);ws=wb.active
            ws.insert_rows(4);ws["D4"]="담당자 확인 메모"
            wb.save(prior);wb.close()
            read=read_prior(prior,period=AccountingPeriod(2026,7))
            self.assertEqual(read.issues,[])
            self.assertEqual(len(read.items),2)
            write_month_end_statement(output,[prior],
                [ReconcileResult(Status.UNPAID,item) for item in read.items],[],[],
                AccountingPeriod(2026,8))
            wb=load_workbook(output)
            try:
                self.assertEqual(wb["26.08"]["D6"].value,"담당자 확인 메모")
            finally:
                wb.close()

    def test_reference_formula_keeps_excel_cached_value_and_original_comment(self):
        from openpyxl.comments import Comment
        with tempfile.TemporaryDirectory() as folder:
            prior=Path(folder)/"prior.xlsx";output=Path(folder)/"result.xlsx"
            write_flat(prior)
            wb=load_workbook(prior);ws=wb.active
            ws["D6"]="원장 잔액";ws["E6"]="=E3+E4+E5"
            ws["E6"].comment=Comment("담당자 확인용","담당자")
            wb.save(prior);wb.close()
            # Model the cached result stored by Excel (openpyxl does not calculate).
            with ZipFile(prior) as archive:
                entries={name:archive.read(name) for name in archive.namelist()}
            xml=ElementTree.fromstring(entries["xl/worksheets/sheet1.xml"])
            ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            cache=xml.find('.//s:c[@r="E6"]/s:v',ns)
            self.assertIsNotNone(cache)
            cache.text="600"
            entries["xl/worksheets/sheet1.xml"]=ElementTree.tostring(xml)
            with ZipFile(prior,"w") as archive:
                for name,content in entries.items():
                    archive.writestr(name,content)
            before=prior.read_bytes()
            read=read_prior(prior,period=AccountingPeriod(2026,7))
            self.assertEqual(read.issues,[])
            write_month_end_statement(output,[prior],
                [ReconcileResult(Status.UNPAID,item) for item in read.items],[],[],
                AccountingPeriod(2026,8))
            wb=load_workbook(output)
            try:
                cell=wb["26.08"]["E6"]
                self.assertEqual(cell.value,600)
                self.assertEqual(cell.data_type,"n")
                self.assertIn("담당자 확인용",cell.comment.text)
                self.assertIn("전월 원본 참고 수식: =E3+E4+E5",cell.comment.text)
            finally:
                wb.close()
            self.assertEqual(prior.read_bytes(),before)

    def test_dated_or_identified_broken_transactions_remain_errors(self):
        with tempfile.TemporaryDirectory() as folder:
            for code,name,when in ((None,None,"2026-07-31"),("000011","가상A",None)):
                with self.subTest(code=code):
                    prior=Path(folder)/"prior.xlsx"
                    write_grouped(prior)
                    wb=load_workbook(prior);ws=wb.active
                    ws.append([code,name,when,"추가 거래",None])
                    wb.save(prior);wb.close()
                    read=read_prior(prior,period=AccountingPeriod(2026,7))
                    self.assertTrue(read.issues)
                    with self.assertRaisesRegex(ValueError,"6행"):
                        write_month_end_statement(Path(folder)/"result.xlsx",[prior],[],[],[],
                                                  AccountingPeriod(2026,8))
