import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from adapters.excel.writer import write_result
from application.service import run_reconciliation
from domain.period import AccountingPeriod
from tests.synthetic_adversarial_factory import build


class FinalOutputGoldenE2ETests(unittest.TestCase):
    """Golden test for the complete path: input Excel -> service -> final result Excel."""

    def test_adversarial_inputs_produce_exact_conservative_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            data=build(tmp)
            period=AccountingPeriod(2026,7)

            run=run_reconciliation(
                [data["prior_path"]],
                [data["raw_path"]],
                {"25301"},
                period,
            )

            out=Path(tmp)/"명세서_대사결과.xlsx"
            write_result(
                out,
                run.results,
                run.new_items,
                run.issues,
                [data["prior_path"],data["raw_path"]],
                period.label,
            )

            wb=load_workbook(out,data_only=True)
            self.assertEqual(
                wb.sheetnames,
                ["요약","확인필요","입력데이터확인","차월명세서 초안","당월신규명세","자동대사완료"],
            )

            # Summary is a fixed golden answer for this adversarial fixture.
            summary=dict(wb["요약"].iter_rows(min_row=1,max_col=2,values_only=True))
            self.assertEqual(summary["대상 회계월"],"2026년 7월")
            self.assertEqual(summary["명세서 대사 대상"],7)
            self.assertEqual(summary["자동 대사 완료"],3)
            self.assertEqual(summary["검토 필요"],5)
            self.assertEqual(summary["대사 검토 필요"],4)
            self.assertEqual(summary["Raw 신규 대변 검토"],1)
            self.assertEqual(summary["입력 데이터 확인"],3)
            self.assertEqual(summary["차월 이월 초안"],0)
            self.assertEqual(summary["차월 이월 초안 금액"],0)
            self.assertEqual(summary["초안 제외 검토건"],5)
            self.assertEqual(summary["대사 기준"],"같은 회계월 명세서 ↔ 더존 Raw / 거래처코드+금액 중심")

            # Unique code+amount matches complete automatically; name/description differences are reference notes.
            completed=list(wb["자동대사완료"].iter_rows(min_row=2,values_only=True))
            self.assertEqual(len(completed),3)
            completed_by_code={row[1]:row for row in completed}
            self.assertEqual(completed_by_code["001001"][4],110000)
            self.assertIn("적요 차이",completed_by_code["001002"][6])
            self.assertTrue(
                "거래처명 차이" in completed_by_code["001004"][6]
                or "거래처명 오타 의심" in completed_by_code["001004"][6]
            )

            # Only code/amount/ambiguity/input-safety exceptions remain for human review.
            review=list(wb["확인필요"].iter_rows(min_row=2,values_only=True))
            self.assertEqual(len(review),4)
            review_headers=[cell.value for cell in wb["확인필요"][1]]
            self.assertNotIn("담당자",review_headers)
            review_codes={row[5] for row in review}
            self.assertEqual(review_codes,{"001003","001005","001006","001007"})

            # Dirty source rows are quarantined and never disappear silently.
            input_issues=list(wb["입력데이터확인"].iter_rows(min_row=2,values_only=True))
            self.assertEqual(len(input_issues),3)
            self.assertEqual(sum(row[0]=="명세서" for row in input_issues),1)
            self.assertEqual(sum(row[0]=="더존" for row in input_issues),2)

            # Raw input issues mean absence of a debit is not trustworthy enough
            # to auto-carry an item into the next-month draft.
            draft_rows=list(wb["차월명세서 초안"].iter_rows(min_row=2,values_only=True))
            self.assertEqual(draft_rows,[])
            unpaid_review=[row for row in review if row[5]=="001006"]
            self.assertEqual(len(unpaid_review),1)
            self.assertEqual(unpaid_review[0][0],"더존 입력 확인 필요")
            self.assertIn("지급 여부를 확정할 수 없음",unpaid_review[0][1])

            # Current-month new item is visible but explicitly withheld from the draft.
            current=list(wb["당월신규명세"].iter_rows(min_row=2,values_only=True))
            self.assertEqual(len(current),1)
            self.assertEqual(current[0][0],"사람 확인 필요")
            self.assertEqual(current[0][3],"002001")
            self.assertEqual(current[0][5],"7월 신규 유류비")
            self.assertEqual(current[0][6],880000)
            self.assertIn("차월 초안 자동포함 안 함",current[0][7])

            # Sheet-tab colors mirror each sheet's visual meaning.
            self.assertTrue(wb["확인필요"].sheet_properties.tabColor.rgb.endswith("FFD966"))
            self.assertTrue(wb["입력데이터확인"].sheet_properties.tabColor.rgb.endswith("E74C3C"))
            self.assertTrue(wb["당월신규명세"].sheet_properties.tabColor.rgb.endswith("8E44AD"))
            self.assertTrue(wb["요약"].sheet_properties.tabColor.rgb.endswith("2F75B5"))
            self.assertTrue(wb["차월명세서 초안"].sheet_properties.tabColor.rgb.endswith("A5A5A5"))
            completed_sheet=wb["자동대사완료"]
            self.assertIsNotNone(completed_sheet.sheet_properties.tabColor)
            self.assertTrue(completed_sheet.sheet_properties.tabColor.rgb.endswith("70AD47"))
            self.assertGreater(completed_sheet.column_dimensions["D"].width,20)
            self.assertGreaterEqual(completed_sheet.row_dimensions[2].height,20)

            wb.close()


if __name__=="__main__":
    unittest.main()
