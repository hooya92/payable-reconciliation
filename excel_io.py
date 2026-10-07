from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from models import JournalLine, PayableItem, ReconcileResult, Status
from reconciliation import normalize_code


DOUZONE_ALIASES = {
    "date": ("기표일자", "일자", "날짜"),
    "account_code": ("계정코드", "계정 코드"),
    "account_name": ("계정과목명", "계정과목", "계정명"),
    "vendor_code": ("거래처코드", "거래처 코드", "코드"),
    "vendor_name": ("거래처명", "거래처 명"),
    "description": ("적요", "적요명"),
    "debit": ("차변", "차변금액", "차변 금액"),
    "credit": ("대변", "대변금액", "대변 금액"),
}
PRIOR_ALIASES = {
    "vendor_code": ("거래처코드", "거래처 코드", "코드"),
    "vendor_name": ("거래처명", "거래처", "업체명"),
    "date": ("날짜", "일자", "기표일자"),
    "description": ("적요", "내역", "내용"),
    "amount": ("금액", "미지급금", "잔액"),
}


def _text(v: object) -> str:
    return "" if v is None else str(v).strip()


def _amount(v: object) -> Decimal:
    if v in (None, ""):
        return Decimal("0")
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v).replace(",", "").replace("₩", "").strip()
    if not s:
        return Decimal("0")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("0")


def _date(v: object) -> str:
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return _text(v)


def _norm_header(v: object) -> str:
    return "".join(_text(v).lower().split())


def _find_header(ws, aliases: dict[str, tuple[str, ...]], required: set[str], max_rows: int = 30):
    for row in range(1, min(ws.max_row, max_rows) + 1):
        values = [_norm_header(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        mapping = {}
        for field, names in aliases.items():
            normalized = {_norm_header(x) for x in names}
            for idx, value in enumerate(values, 1):
                if value in normalized:
                    mapping[field] = idx
                    break
        if required.issubset(mapping):
            return row, mapping
    raise ValueError("필요한 컬럼을 찾지 못했습니다. 헤더명이 예상 구조와 다른지 확인해주세요.")


def read_douzone(path: str | Path, account_codes: set[str] | None = None) -> list[JournalLine]:
    wb = load_workbook(path, read_only=True, data_only=True)
    result: list[JournalLine] = []
    wanted = {normalize_code(x) for x in account_codes or set() if x}
    for ws in wb.worksheets:
        try:
            header, cols = _find_header(ws, DOUZONE_ALIASES,
                {"vendor_code", "vendor_name", "description", "debit", "credit"})
        except ValueError:
            continue
        for r in range(header + 1, ws.max_row + 1):
            account_code = normalize_code(ws.cell(r, cols["account_code"]).value) if "account_code" in cols else ""
            account_name = _text(ws.cell(r, cols["account_name"]).value) if "account_name" in cols else ""
            if wanted and account_code not in wanted:
                continue
            if not wanted and account_name and "미지급" not in account_name:
                continue
            debit = _amount(ws.cell(r, cols["debit"]).value)
            credit = _amount(ws.cell(r, cols["credit"]).value)
            if debit == 0 and credit == 0:
                continue
            result.append(JournalLine(
                vendor_code=normalize_code(ws.cell(r, cols["vendor_code"]).value),
                vendor_name=_text(ws.cell(r, cols["vendor_name"]).value),
                account_code=account_code,
                account_name=account_name,
                description=_text(ws.cell(r, cols["description"]).value),
                debit=debit, credit=credit,
                date=_date(ws.cell(r, cols["date"]).value) if "date" in cols else "",
                row_number=r,
            ))
    wb.close()
    if not result:
        raise ValueError("더존 파일에서 대상 미지급금 전표를 찾지 못했습니다.")
    return result


def read_prior(path: str | Path) -> list[PayableItem]:
    wb = load_workbook(path, read_only=True, data_only=True)
    result: list[PayableItem] = []
    for ws in wb.worksheets:
        try:
            header, cols = _find_header(ws, PRIOR_ALIASES,
                {"vendor_code", "vendor_name", "description", "amount"})
        except ValueError:
            continue
        for r in range(header + 1, ws.max_row + 1):
            code = normalize_code(ws.cell(r, cols["vendor_code"]).value)
            amount = _amount(ws.cell(r, cols["amount"]).value)
            desc = _text(ws.cell(r, cols["description"]).value)
            if not code or amount <= 0 or not desc:
                continue
            result.append(PayableItem(
                vendor_code=code,
                vendor_name=_text(ws.cell(r, cols["vendor_name"]).value),
                description=desc,
                amount=amount,
                date=_date(ws.cell(r, cols["date"]).value) if "date" in cols else "",
                row_number=r,
            ))
    wb.close()
    if not result:
        raise ValueError("전월 명세서에서 대사할 항목을 찾지 못했습니다.")
    return result


def write_result(path: str | Path, results: list[ReconcileResult], new_items: list[JournalLine],
                 source_paths: Iterable[str | Path]) -> None:
    out = Path(path).resolve()
    if any(out == Path(p).resolve() for p in source_paths):
        raise ValueError("원본 Excel에는 저장할 수 없습니다. 다른 파일명을 선택해주세요.")

    wb = Workbook()
    ws = wb.active
    ws.title = "확인필요"
    headers = ["상태", "사유", "거래처코드", "거래처명", "전월적요", "전월금액",
               "더존거래처코드", "더존거래처명", "더존적요", "더존차변", "더존행"]
    ws.append(headers)
    for x in results:
        if x.status == Status.MATCHED:
            continue
        j = x.journal
        ws.append([x.status.value, x.reason, x.prior.vendor_code, x.prior.vendor_name,
                   x.prior.description, float(x.prior.amount),
                   j.vendor_code if j else "", j.vendor_name if j else "",
                   j.description if j else "", float(j.debit) if j else "", j.row_number if j else ""])

    matched = wb.create_sheet("자동대사완료")
    matched.append(headers)
    for x in results:
        if x.status != Status.MATCHED:
            continue
        j = x.journal
        matched.append([x.status.value, x.reason, x.prior.vendor_code, x.prior.vendor_name,
                        x.prior.description, float(x.prior.amount),
                        j.vendor_code if j else "", j.vendor_name if j else "",
                        j.description if j else "", float(j.debit) if j else "", j.row_number if j else ""])

    new_ws = wb.create_sheet("신규미지급")
    new_ws.append(["기표일자", "계정코드", "계정과목명", "거래처코드", "거래처명", "적요", "대변"])
    for j in new_items:
        new_ws.append([j.date, j.account_code, j.account_name, j.vendor_code, j.vendor_name,
                       j.description, float(j.credit)])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E9EEF5")
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            width = min(max((len(str(c.value or "")) for c in col), default=8) + 2, 42)
            sheet.column_dimensions[col[0].column_letter].width = width

    wb.save(out)
