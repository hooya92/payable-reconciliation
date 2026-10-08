from pathlib import Path
from typing import Iterable
from openpyxl import Workbook
from domain.models import Status
from adapters.excel.styles import style_workbook


REVIEW_TAB_COLOR="FFD60A"


def _mark_review_tab(sheet, has_items):
    if has_items:
        sheet.sheet_properties.tabColor=REVIEW_TAB_COLOR


def write_result(path, results, new_items, issues, source_paths:Iterable, period_label:str):
    out=Path(path).resolve()
    if any(out==Path(p).resolve() for p in source_paths):
        raise ValueError("원본 Excel에는 저장할 수 없습니다.")
    wb=Workbook(); ws=wb.active; ws.title="확인필요"
    ws.append(["상태","사유","원본파일","원본시트","원본행","거래처코드","거래처명","명세서적요","명세서금액","더존거래처","더존적요","더존차변","더존행"])
    review_results=[x for x in results if x.status not in (Status.MATCHED, Status.PARTIAL)]
    for x in review_results:
        j=x.journal
        ws.append([x.status.value,x.reason,x.prior.source.file_name,x.prior.source.sheet,x.prior.source.row,x.prior.vendor_code,x.prior.vendor_name,x.prior.description,int(x.prior.amount), j.vendor_name if j else "",j.description if j else "",int(j.debit) if j else "",j.row_number if j else ""])
    _mark_review_tab(ws,bool(review_results))
    iq=wb.create_sheet("입력데이터확인"); iq.append(["출처","시트","행","필드","원본값","사유"])
    for x in issues: iq.append([x.source,x.sheet,x.row,x.field,str(x.raw_value),x.reason])
    _mark_review_tab(iq,bool(issues))
    # Only unambiguous outstanding opening balances and current credits are carried.
    draft=wb.create_sheet("당월말명세서 초안")
    draft.append(["구분","거래처코드","거래처명","날짜","적요","금액","근거","원본"])
    draft_rows=[]
    for x in results:
        if x.status in (Status.UNPAID, Status.PARTIAL):
            p=x.prior
            balance = p.amount - x.journal.debit if x.status == Status.PARTIAL else p.amount
            label = "부분지급잔액" if x.status == Status.PARTIAL else "미지급이월"
            basis = x.reason if x.status == Status.PARTIAL else "당월 더존 대응 차변 없음"
            draft_rows.append([label,p.vendor_code,p.vendor_name,p.date,p.description,int(balance),
                               basis,f"{p.source.file_name} · {p.source.row}행"])

    for item in new_items:
        if not item.auto_carry or item.remaining <= 0:
            continue
        j=item.journal
        draft_rows.append(["당월신규",j.vendor_code,j.vendor_name,j.date,j.description,int(item.remaining),
                           item.reason,f"더존 Raw · {j.row_number}행"])
    draft_rows.sort(key=lambda x:(x[2] or "",x[1] or "",x[3] or "",x[4] or ""))

    current=None; subtotal=0
    for row in draft_rows:
        vendor=(row[1],row[2])
        if current is not None and vendor != current:
            draft.append(["소계",current[0],current[1],"","",subtotal,"거래처 소계",""])
            subtotal=0
        draft.append(row); subtotal += row[5]; current=vendor
    if current is not None:
        draft.append(["소계",current[0],current[1],"","",subtotal,"거래처 소계",""])
    if draft_rows:
        draft.append(["전체합계","","","","",sum(r[5] for r in draft_rows),"당월말 명세서 초안 합계",""])


    nw=wb.create_sheet("당월신규명세"); nw.append(["검토상태","기표일자","계정코드","거래처코드","거래처명","적요","대변","사유"])
    for item in new_items:
        j=item.journal
        state = ("초안 반영 · 최종 확인" if item.remaining > 0 else "당월 전액 상계") if item.auto_carry else "사람 확인 필요"
        note = f"{item.reason} · 당월말 잔액 {item.remaining:,.0f}원" if item.auto_carry else item.reason
        nw.append([state,j.date,j.account_code,j.vendor_code,j.vendor_name,j.description,int(j.credit),note])
    _mark_review_tab(nw,bool(new_items))
    ok=wb.create_sheet("자동대사완료"); ok.append(["상태","거래처코드","거래처명","명세서적요","금액","더존행","참고"])
    for x in results:
        if x.status in (Status.MATCHED, Status.PARTIAL):
            note=x.reason if x.status==Status.PARTIAL or x.rule=="CODE_AMOUNT_UNIQUE_WITH_NOTE" else ""
            ok.append([x.status.value,x.prior.vendor_code,x.prior.vendor_name,x.prior.description,int(x.prior.amount),x.journal.row_number,note])
    info=wb.create_sheet("요약",0)
    matched_count=sum(x.status in (Status.MATCHED, Status.PARTIAL) for x in results)
    reconciliation_review=sum(x.status not in (Status.MATCHED, Status.PARTIAL) for x in results)
    total_review=reconciliation_review+len(new_items)
    info.append(["대상 회계월",period_label])
    info.append(["명세서 대사 대상",len(results)])
    info.append(["자동 대사 완료",matched_count])
    info.append(["검토 필요",total_review])
    info.append(["대사 검토 필요",reconciliation_review])
    info.append(["Raw 신규 대변 검토",len(new_items)])
    info.append(["입력 데이터 확인",len(issues)])
    info.append(["당월말 명세 초안",len(draft_rows)])
    info.append(["당월말 명세 금액",sum(r[5] for r in draft_rows)])
    info.append(["초안 제외 검토건",sum(x.status not in (Status.MATCHED,Status.PARTIAL,Status.UNPAID) for x in results)+sum(not x.auto_carry for x in new_items)])
    info.append(["대사 기준","전월 말 명세서 ↔ 당월 더존 Raw / 거래처코드+금액·적요 중심"])
    info.append(["초안 원칙","확실한 전월 잔액·부분지급 잔액·당월 신규 잔액만 반영 / 나머지는 확인 필요"])

    style_workbook(wb, draft, info)
    wb.save(out)
