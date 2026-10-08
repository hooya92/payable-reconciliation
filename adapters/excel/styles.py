import math
import unicodedata

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


PALETTE={
    "navy":"17365D",
    "blue":"2F75B5",
    "blue_light":"EAF3FB",
    "green":"70AD47",
    "green_light":"E2F0D9",
    "yellow":"FFD966",
    "yellow_light":"FFF4CC",
    "orange_light":"FCE4D6",
    "red_light":"FDE9E7",
    "purple_light":"EDE7F6",
    "gray":"F3F5F7",
    "gray_dark":"D9E1F2",
    "white":"FFFFFF",
    "text":"1F2937",
    "muted":"6B7280",
    "subtotal":"E7E6E6",
    "total":"D9E2F3",
}
AMOUNT_HEADERS={"금액","명세서금액","더존차변","대변","차월 초안 금액","차월 이월 초안 금액"}
_CENTER_HEADERS={"상태","검토상태","출처","시트","행","필드","원본시트","원본행","거래처코드","계정코드","더존행","구분","기표일자","날짜"}
_THIN=Side(style="thin",color="D9DEE7")
_MEDIUM=Side(style="medium",color="B8C2D1")


def _display_width(value):
    if value is None:
        return 0
    return sum(
        2 if unicodedata.east_asian_width(ch) in ("W","F") else 1
        for ch in str(value)
    )


def _autofit_sheet(sheet):
    widths={}
    for col in sheet.columns:
        letter=col[0].column_letter
        header=str(col[0].value or "")
        content_width=max((_display_width(cell.value) for cell in col),default=0)
        if any(key in header for key in ("적요","사유","참고","원본파일","거래처명","원본값","근거")):
            min_width,max_width=24,46
        elif any(key in header for key in ("코드","행","상태","구분","시트","필드")):
            min_width,max_width=10,20
        else:
            min_width,max_width=10,30
        width=min(max(content_width+3,min_width),max_width)
        sheet.column_dimensions[letter].width=width
        widths[letter]=width

    sheet.row_dimensions[1].height=max(sheet.row_dimensions[1].height or 0,28)
    for row_idx in range(2,sheet.max_row+1):
        lines=1
        for cell in sheet[row_idx]:
            if cell.value in (None,""):
                continue
            usable=max(widths.get(cell.column_letter,10)-2,1)
            needed=sum(
                max(1,math.ceil(_display_width(part)/usable))
                for part in str(cell.value).splitlines()
            )
            lines=max(lines,needed)
        sheet.row_dimensions[row_idx].height=min(max(22,lines*18),90)


def _style_header(sheet):
    for cell in sheet[1]:
        cell.font=Font(name="맑은 고딕",size=10,bold=True,color=PALETTE["white"])
        cell.fill=PatternFill("solid",fgColor=PALETTE["navy"])
        cell.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True)
        cell.border=Border(bottom=_MEDIUM)


def _style_body(sheet):
    for rr,row in enumerate(sheet.iter_rows(min_row=2),start=2):
        if rr % 2 == 0:
            zebra=PatternFill("solid",fgColor="FAFBFC")
        else:
            zebra=None
        for cell in row:
            cell.font=Font(name="맑은 고딕",size=10,color=PALETTE["text"])
            cell.alignment=Alignment(vertical="top",wrap_text=True)
            cell.border=Border(bottom=_THIN)
            if zebra and cell.fill.fill_type is None:
                cell.fill=zebra


def _apply_column_alignment(sheet):
    for col_idx,cell in enumerate(sheet[1],1):
        header=str(cell.value or "")
        if header in AMOUNT_HEADERS:
            for rr in range(2,sheet.max_row+1):
                c=sheet.cell(rr,col_idx)
                c.number_format="#,##0"
                c.alignment=Alignment(horizontal="right",vertical="top",wrap_text=True)
        elif header in _CENTER_HEADERS:
            for rr in range(2,sheet.max_row+1):
                sheet.cell(rr,col_idx).alignment=Alignment(horizontal="center",vertical="top",wrap_text=True)


def _fill_row(sheet, row_index, color, bold=False):
    fill=PatternFill("solid",fgColor=color)
    for cell in sheet[row_index]:
        cell.fill=fill
        if bold:
            cell.font=Font(name="맑은 고딕",size=10,bold=True,color=PALETTE["text"])


def _style_review_sheet(sheet):
    for rr in range(2,sheet.max_row+1):
        status=str(sheet.cell(rr,1).value or "")
        if "입력 확인" in status:
            _fill_row(sheet,rr,PALETTE["red_light"])
        elif "거래처 오류" in status:
            _fill_row(sheet,rr,PALETTE["orange_light"])
        elif status:
            _fill_row(sheet,rr,PALETTE["yellow_light"])


def _style_issue_sheet(sheet):
    for rr in range(2,sheet.max_row+1):
        _fill_row(sheet,rr,"FFF9E6")


def _style_new_sheet(sheet):
    for rr in range(2,sheet.max_row+1):
        _fill_row(sheet,rr,PALETTE["purple_light"])


def _style_completed_sheet(sheet):
    for rr in range(2,sheet.max_row+1):
        _fill_row(sheet,rr,PALETTE["green_light"])


def _style_summary(info):
    info.sheet_view.showGridLines=False
    info.freeze_panes=None
    info.auto_filter.ref=None
    info.sheet_properties.tabColor=PALETTE["blue"]

    info.column_dimensions["A"].width=24
    info.column_dimensions["B"].width=60
    for rr in range(1,info.max_row+1):
        label=info.cell(rr,1)
        value=info.cell(rr,2)
        label.font=Font(name="맑은 고딕",size=10,bold=True,color=PALETTE["text"])
        label.fill=PatternFill("solid",fgColor=PALETTE["blue_light"])
        label.alignment=Alignment(vertical="center",horizontal="left")
        value.font=Font(name="맑은 고딕",size=10,color=PALETTE["text"])
        value.alignment=Alignment(vertical="center",horizontal="left",wrap_text=True)
        label.border=Border(bottom=_THIN)
        value.border=Border(bottom=_THIN)
        info.row_dimensions[rr].height=max(info.row_dimensions[rr].height or 0,24)

    # KPI rows get stronger emphasis.
    highlight={
        "자동 대사 완료":PALETTE["green_light"],
        "검토 필요":PALETTE["yellow_light"],
        "입력 데이터 확인":PALETTE["red_light"],
        "Raw 신규 대변 검토":PALETTE["purple_light"],
    }
    for rr in range(1,info.max_row+1):
        label=str(info.cell(rr,1).value or "")
        if label in highlight:
            info.cell(rr,2).fill=PatternFill("solid",fgColor=highlight[label])
            info.cell(rr,2).font=Font(name="맑은 고딕",size=12,bold=True,color=PALETTE["text"])
        if "금액" in label:
            info.cell(rr,2).number_format="#,##0"


def style_workbook(wb, draft, info):
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines=False
        sheet.sheet_view.zoomScale=90
        _style_header(sheet)
        _style_body(sheet)
        _apply_column_alignment(sheet)

        sheet.freeze_panes="A2"
        sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.fitToWidth=1
        sheet.page_setup.fitToHeight=0
        sheet.page_margins.left=0.25
        sheet.page_margins.right=0.25
        sheet.page_margins.top=0.5
        sheet.page_margins.bottom=0.5
        sheet.oddFooter.center.text="페이지 &P / &N"
        _autofit_sheet(sheet)

    # Sheet-specific colors and row emphasis.
    if "확인필요" in wb.sheetnames:
        review=wb["확인필요"]
        review.sheet_properties.tabColor=PALETTE["yellow"]
        _style_review_sheet(review)

    if "입력데이터확인" in wb.sheetnames:
        issue=wb["입력데이터확인"]
        issue.sheet_properties.tabColor="E74C3C"
        _style_issue_sheet(issue)

    if "당월신규명세" in wb.sheetnames:
        new_sheet=wb["당월신규명세"]
        new_sheet.sheet_properties.tabColor="8E44AD"
        _style_new_sheet(new_sheet)

    if "자동대사완료" in wb.sheetnames:
        completed=wb["자동대사완료"]
        completed.sheet_properties.tabColor=PALETTE["green"]
        _style_completed_sheet(completed)

    if "당월말명세서 초안" in wb.sheetnames:
        draft.sheet_properties.tabColor="A5A5A5"

    for rr in range(2,draft.max_row+1):
        kind=draft.cell(rr,1).value
        fill={
            "전월이월":PALETTE["yellow_light"],
            "미지급이월":PALETTE["yellow_light"],
            "부분지급잔액":PALETTE["yellow_light"],
            "당월신규":PALETTE["purple_light"],
            "소계":PALETTE["subtotal"],
            "전체합계":PALETTE["total"],
        }.get(kind)
        if fill:
            _fill_row(draft,rr,fill,bold=kind in ("소계","전체합계"))

    draft.page_setup.orientation="landscape"
    draft.print_title_rows="1:1"
    draft.column_dimensions["E"].width=max(draft.column_dimensions["E"].width or 0,36)
    draft.column_dimensions["G"].width=max(draft.column_dimensions["G"].width or 0,22)
    draft.column_dimensions["H"].width=max(draft.column_dimensions["H"].width or 0,28)

    _style_summary(info)
