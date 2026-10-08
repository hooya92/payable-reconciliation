from __future__ import annotations
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import re
import struct
from itertools import islice

try:
    import xlrd
    from xlrd import compdoc
except ImportError:
    xlrd=None
    compdoc=None

from openpyxl import load_workbook as _openpyxl_load_workbook
from openpyxl.utils import get_column_letter

from domain.models import DataQuality, JournalLine, PayableItem, SourceRef
from domain.normalization import normalize_code, parse_amount
from domain.period import AccountingPeriod, parse_date


class _XlsCell:
    def __init__(self, book, sheet, rowx, colx):
        if rowx >= sheet.nrows or colx >= sheet.ncols:
            self.value=None
            self.data_type="n"
            self.number_format="General"
            return

        cell=sheet.cell(rowx,colx)
        value=cell.value
        if cell.ctype==xlrd.XL_CELL_DATE:
            value=xlrd.xldate_as_datetime(value,book.datemode)
        elif cell.ctype==xlrd.XL_CELL_BOOLEAN:
            value=bool(value)
        elif cell.ctype in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK):
            value=None

        self.value=value
        self.data_type="s" if cell.ctype==xlrd.XL_CELL_TEXT else "n"
        self.number_format="General"
        if cell.xf_index is not None and 0 <= cell.xf_index < len(book.xf_list):
            xf=book.xf_list[cell.xf_index]
            fmt=book.format_map.get(xf.format_key)
            if fmt is not None:
                self.number_format=fmt.format_str or "General"


class _XlsSheet:
    def __init__(self, book, sheet):
        self._book=book
        self._sheet=sheet
        self.title=sheet.name
        self.max_row=sheet.nrows
        self.max_column=sheet.ncols

    def cell(self, row, column):
        return _XlsCell(self._book,self._sheet,row-1,column-1)

    def iter_rows(self, min_row=1, values_only=False):
        for rowx in range(max(min_row-1,0),self.max_row):
            cells=tuple(self.cell(rowx+1,colx+1) for colx in range(self.max_column))
            if values_only:
                yield tuple(cell.value for cell in cells)
            else:
                yield cells


class _XlsWorkbook:
    def __init__(self, book):
        self._book=book
        self.worksheets=[_XlsSheet(book,book.sheet_by_index(i)) for i in range(book.nsheets)]

    def close(self):
        self._book.release_resources()


def _xls_contains_formula(path):
    if xlrd is None or compdoc is None:
        raise ValueError(
            "구형 .xls 파일 지원 모듈(xlrd)이 설치되어 있지 않습니다. "
            "프로젝트 폴더에서 'python -m pip install -r requirements.txt'를 실행해주세요."
        )
    raw=Path(path).read_bytes()
    try:
        if raw[:8]==compdoc.SIGNATURE:
            doc=compdoc.CompDoc(raw)
            stream=doc.get_named_stream("Workbook") or doc.get_named_stream("Book")
        else:
            stream=raw
    except Exception as exc:
        raise ValueError(f"구형 .xls 파일 구조를 읽을 수 없습니다: {exc}") from exc

    if not stream:
        raise ValueError("구형 .xls 파일에서 Workbook 스트림을 찾지 못했습니다.")

    pos=0
    while pos+4 <= len(stream):
        opcode,length=struct.unpack("<HH",stream[pos:pos+4])
        end=pos+4+length
        if end>len(stream):
            raise ValueError("구형 .xls 파일의 BIFF 레코드가 손상되어 안전하게 읽을 수 없습니다.")
        if opcode==0x0006:  # BIFF FORMULA record
            return True
        pos=end
    return False


def _load_workbook(path, **kwargs):
    path=Path(path)
    if path.suffix.lower()==".xls":
        if xlrd is None:
            raise ValueError(
                "구형 .xls 파일 지원 모듈(xlrd)이 설치되어 있지 않습니다. "
                "프로젝트 폴더에서 'python -m pip install -r requirements.txt'를 실행해주세요."
            )
        if _xls_contains_formula(path):
            raise ValueError(
                "구형 .xls 파일에 수식 셀이 포함되어 있습니다. "
                "계산 캐시값을 정상 데이터로 오인하지 않도록 자동 처리를 중단합니다. "
                "Excel에서 값이 확정된 .xlsx로 다시 저장한 뒤 선택해주세요."
            )
        try:
            book=xlrd.open_workbook(str(path),formatting_info=True,on_demand=False)
        except Exception as exc:
            raise ValueError(f"구형 .xls 파일을 읽을 수 없습니다: {exc}") from exc
        return _XlsWorkbook(book)
    return _openpyxl_load_workbook(path,**kwargs)


@dataclass(frozen=True)
class InputIssue:
    source: str
    sheet: str
    row: int
    field: str
    raw_value: object
    reason: str


@dataclass
class ReadResult:
    items: list = field(default_factory=list)
    issues: list[InputIssue] = field(default_factory=list)
    detected_headers: list[str] = field(default_factory=list)
    recognized_sheets: list[str] = field(default_factory=list)
    # Signed subtotal groups carry a verified vendor balance, but the negative
    # entries cannot safely be assigned to individual invoices automatically.
    review_items: list[tuple[PayableItem, str]] = field(default_factory=list)


def _text(v): return "" if v is None else str(v).strip()
def _is_formula(cell): return getattr(cell,"data_type",None)=="f" or (isinstance(cell.value,str) and cell.value.startswith("="))
def _code_from_cell(cell):
    """Preserve identifier leading zeroes when Excel stores a numeric code with a zero-only display format."""
    value=cell.value
    if value is None: return ""
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        fmt=(cell.number_format or "").strip()
        if re.fullmatch(r"0+",fmt):
            if isinstance(value,float) and not value.is_integer():
                return normalize_code(value)
            return str(int(value)).zfill(len(fmt))
    return normalize_code(value)
def _header(v): return "".join(_text(v).lower().split())
def _date(v):
    if isinstance(v, (datetime, date)): return v.strftime("%Y-%m-%d")
    return _text(v)
def _is_summary_label(v):
    return _header(v) in {"소계","합계","총계","전체합계"}


# Conservative aliases observed in common Douzone/iCUBE-style journal exports.
# File names are intentionally ignored; classification is based on worksheet structure.
DOUZONE_DATE_HEADERS=("기표일자","기표일","전표일자","전표일","회계일자","일자","날짜")
DOUZONE_ACCOUNT_CODE_HEADERS=("계정코드","계정 코드","계정과목코드","계정과목 코드")
DOUZONE_ACCOUNT_NAME_HEADERS=("계정과목명","계정과목","계정명")
DOUZONE_VENDOR_CODE_HEADERS=("거래처코드","거래처 코드","거래처번호","거래처 번호","거래처No","거래처 No")
DOUZONE_VENDOR_NAME_HEADERS=("거래처명","거래처 명","거래처","업체명")
DOUZONE_DESCRIPTION_HEADERS=("적요","적요명","전표적요","전표 적요","내용")
DOUZONE_DEBIT_HEADERS=("차변","차변금액","차변 금액","차변액")
DOUZONE_CREDIT_HEADERS=("대변","대변금액","대변 금액","대변액")
DOUZONE_REQUIRED_GROUPS=(
    DOUZONE_VENDOR_NAME_HEADERS,
    DOUZONE_DESCRIPTION_HEADERS,
    DOUZONE_DEBIT_HEADERS,
    DOUZONE_CREDIT_HEADERS,
)


def _period_from_sheet_title(title):
    """Parse common monthly sheet names such as 26.07, 2026-07, 2026년 7월."""
    text=_text(title).strip()
    patterns=(
        r"(?<!\d)(20\d{2})[.\-_/](0?[1-9]|1[0-2])(?!\d)",
        r"(?<!\d)(\d{2})[.\-_/](0?[1-9]|1[0-2])(?!\d)",
        r"(?<!\d)(20\d{2})\s*년\s*(0?[1-9]|1[0-2])\s*월",
        r"(?<!\d)(\d{2})\s*년\s*(0?[1-9]|1[0-2])\s*월",
    )
    for pattern in patterns:
        m=re.search(pattern,text,re.IGNORECASE)
        if not m:
            continue
        year=int(m.group(1)); month=int(m.group(2))
        if year < 100:
            year += 2000
        return AccountingPeriod(year,month)
    return None


def _sheet_matches_period(title, period):
    return period is not None and _period_from_sheet_title(title)==period


def detect_statement_periods(path: str|Path) -> list[AccountingPeriod]:
    """Return month periods from statement-like worksheet titles only."""
    wb=_load_workbook(path,read_only=True,data_only=False)
    groups=[("거래처코드","거래처 코드","코드"),("거래처명","거래처","업체명"),("적요","내역","내용"),("금액","미지급금","잔액")]
    found=set()
    try:
        for ws in wb.worksheets:
            hr,_vals=_header_row(ws,groups)
            if not hr:
                continue
            period=_period_from_sheet_title(ws.title)
            if period:
                found.add(period)
    finally:
        wb.close()
    return sorted(found)


def _select_prior_sheets(wb, groups, period):
    """Pick only the intended monthly statement sheet when a workbook contains many month tabs."""
    full=[]
    for ws in wb.worksheets:
        hr,vals=_header_row(ws,groups)
        if hr:
            full.append((ws,hr,vals))

    if period is not None:
        title_matches=[entry for entry in full if _sheet_matches_period(entry[0].title,period)]
        if len(title_matches)==1:
            return title_matches
        if len(title_matches)>1:
            names=", ".join(x[0].title for x in title_matches)
            raise ValueError(f"{period.label} 명세서 시트가 여러 개라 자동 선택할 수 없습니다: {names}")
        if len(full)>1:
            names=", ".join(x[0].title for x in full)
            found=sorted({
                p for p in (_period_from_sheet_title(x[0].title) for x in full) if p
            })
            found_text=", ".join(p.label for p in found) if found else names
            raise ValueError(
                f"현재 대상 회계월은 {period.label}이며 같은 월 명세서가 필요합니다. "
                f"선택한 파일에서 확인된 명세서 월: {found_text}. "
                f"{period.label} 시트가 포함된 파일을 선택하거나 대상 회계월을 변경해주세요."
            )

    if len(full)==1:
        detected = _period_from_sheet_title(full[0][0].title)
        if period is not None and detected is not None and detected != period:
            raise ValueError(
                f"전월 명세서는 {period.label}이어야 하지만 선택된 시트는 {detected.label}입니다."
            )
        return full
    if len(full)>1:
        names=", ".join(x[0].title for x in full)
        raise ValueError(f"명세서 데이터 시트가 여러 개라 자동 선택할 수 없습니다: {names}")

    # No full header was found. A data-like partial header must not be silently ignored.
    for ws in wb.worksheets:
        partial_row,matched=_partial_header_row(ws,groups,3)
        if partial_row:
            raise ValueError(
                f"명세서의 '{ws.title}' 시트가 데이터 표처럼 보이지만 필수 헤더를 완전히 식별하지 못했습니다 "
                f"(필수 항목 {matched}/4개 인식). 시트 구조를 확인해주세요."
            )
    return []


def _header_row(ws, required_groups, max_rows=30):
    # Stream rows once; repeated ws.cell() calls reparse read-only .xlsx sheets.
    # Some valid exporters omit the worksheet <dimension> tag. Streaming works
    # even when the read-only worksheet cannot know its row count up front.
    limit=min(ws.max_row, max_rows) if ws.max_row is not None else max_rows
    for r, row in enumerate(islice(ws.iter_rows(min_row=1, values_only=True), limit), start=1):
        vals=[_header(value) for value in row]
        if all(any(name in vals for name in group) for group in required_groups):
            return r, vals
    return None, []


def _partial_header_row(ws, required_groups, min_groups, max_rows=30):
    """Detect a likely data table whose headers changed enough that we should not silently ignore it."""
    best=(None,0)
    # Some valid exporters omit the worksheet <dimension> tag. Streaming works
    # even when the read-only worksheet cannot know its row count up front.
    limit=min(ws.max_row, max_rows) if ws.max_row is not None else max_rows
    for r, row in enumerate(islice(ws.iter_rows(min_row=1, values_only=True), limit), start=1):
        vals=[_header(value) for value in row]
        matched=sum(any(_header(name) in vals for name in group) for group in required_groups)
        if matched>best[1]:
            best=(r,matched)
    return best if best[1]>=min_groups else (None,0)


def _first(vals, *names):
    for name in names:
        try: return vals.index(_header(name))+1
        except ValueError: pass
    return None


def _douzone_columns(vals):
    # 더존 전표출력은 '코드'가 2개일 수 있다.
    # 계정과목명 바로 왼쪽 코드=계정코드, 거래처명 바로 왼쪽 코드=거래처코드.
    account_name=_first(vals,*DOUZONE_ACCOUNT_NAME_HEADERS)
    vendor_name=_first(vals,*DOUZONE_VENDOR_NAME_HEADERS)
    account_code=_first(vals,*DOUZONE_ACCOUNT_CODE_HEADERS)
    vendor_code=_first(vals,*DOUZONE_VENDOR_CODE_HEADERS)
    if account_code is None and account_name and account_name>1 and vals[account_name-2]=="코드":
        account_code=account_name-1
    if vendor_code is None and vendor_name and vendor_name>1 and vals[vendor_name-2]=="코드":
        vendor_code=vendor_name-1
    return {
        "date":_first(vals,*DOUZONE_DATE_HEADERS),
        "account_code":account_code, "account_name":account_name,
        "vendor_code":vendor_code, "vendor_name":vendor_name,
        "description":_first(vals,*DOUZONE_DESCRIPTION_HEADERS),
        "debit":_first(vals,*DOUZONE_DEBIT_HEADERS),
        "credit":_first(vals,*DOUZONE_CREDIT_HEADERS),
    }


def detect_douzone_periods(path: str|Path, account_codes:set[str]|None=None) -> list[AccountingPeriod]:
    """Detect months that actually contain rows for the selected account code(s) in a Douzone Raw file."""
    wanted={normalize_code(x) for x in (account_codes or set()) if x}
    wb=_load_workbook(path,read_only=True,data_only=False)
    groups=DOUZONE_REQUIRED_GROUPS
    found=set()
    try:
        for ws in wb.worksheets:
            hr,vals=_header_row(ws,groups)
            if not hr:
                continue
            cols=_douzone_columns(vals)
            if not cols["date"] or not cols["account_code"]:
                continue
            for row in ws.iter_rows(min_row=hr+1,values_only=False):
                account_cell=row[cols["account_code"]-1]
                date_cell=row[cols["date"]-1]
                if _is_formula(account_cell) or _is_formula(date_cell):
                    continue
                ac=_code_from_cell(account_cell)
                if wanted and ac not in wanted:
                    continue
                parsed=parse_date(date_cell.value)
                if parsed is not None:
                    found.add(AccountingPeriod(parsed.year,parsed.month))
    finally:
        wb.close()
    return sorted(found)


def classify_excel_input(path: str|Path) -> str:
    """Classify an Excel input structurally as 'prior', 'douzone', 'ambiguous', or 'unknown'."""
    wb=_load_workbook(path,read_only=True,data_only=False)
    prior_groups=[
        ("거래처코드","거래처 코드","코드"),
        ("거래처명","거래처","업체명"),
        ("적요","내역","내용"),
        ("금액","미지급금","잔액"),
    ]
    douzone_groups=DOUZONE_REQUIRED_GROUPS
    prior_found=False
    douzone_found=False
    try:
        for ws in wb.worksheets:
            phr,pvals=_header_row(ws,prior_groups)
            if phr:
                prior_found=True

            dhr,dvals=_header_row(ws,douzone_groups)
            if dhr:
                cols=_douzone_columns(dvals)
                # A real Douzone journal must have both account and vendor identifiers.
                if cols["account_code"] and cols["vendor_code"]:
                    douzone_found=True

            if prior_found and douzone_found:
                return "ambiguous"
    finally:
        wb.close()

    if douzone_found:
        return "douzone"
    if prior_found:
        return "prior"
    return "unknown"


def read_douzone(path: str|Path, account_codes:set[str]|None=None, period: AccountingPeriod|None=None) -> ReadResult:
    wanted={normalize_code(x) for x in (account_codes or set()) if x}
    if not wanted:
        raise ValueError("미지급금 계정코드를 1개 이상 지정해야 합니다.")
    wb=_load_workbook(path,read_only=True,data_only=False)
    out=ReadResult()
    groups=DOUZONE_REQUIRED_GROUPS
    try:
        for ws in wb.worksheets:
            hr, vals=_header_row(ws,groups)
            if not hr:
                partial_row,matched=_partial_header_row(ws,groups,3)
                if partial_row:
                    raise ValueError(
                        f"더존 파일의 '{ws.title}' 시트가 전표 표처럼 보이지만 필수 헤더를 완전히 식별하지 못했습니다 "
                        f"(필수 항목 {matched}/4개 인식). 시트 구조를 확인해주세요."
                    )
                continue
            out.recognized_sheets.append(ws.title)
            cols=_douzone_columns(vals)
            needed=["vendor_code","vendor_name","description","debit","credit","account_code"]
            if period: needed.append("date")
            missing=[x for x in needed if not cols[x]]
            if missing:
                raise ValueError("더존 파일 필수 헤더를 안전하게 식별하지 못했습니다: "+", ".join(missing))
            out.detected_headers=[_text(v) for v in next(
                islice(ws.iter_rows(min_row=hr,values_only=True),1), ()
            )]
            for r, row in enumerate(ws.iter_rows(min_row=hr+1, values_only=False), start=hr+1):
                account_cell=row[cols["account_code"]-1]
                if _is_formula(account_cell):
                    out.issues.append(InputIssue("더존",ws.title,r,"계정코드",account_cell.value,"수식 셀은 대상 계정 여부를 확정할 수 없어 자동 처리하지 않음")); continue
                ac=_code_from_cell(account_cell)
                if ac not in wanted:
                    continue
                an=_text(row[cols["account_name"]-1].value) if cols["account_name"] else ""

                date_cell=row[cols["date"]-1] if cols["date"] else None
                raw_date=date_cell.value if date_cell else None
                if date_cell is not None and _is_formula(date_cell):
                    out.issues.append(InputIssue("더존",ws.title,r,"기표일자",raw_date,"수식 날짜는 대상월을 확정할 수 없어 자동 처리하지 않음")); continue
                if period:
                    parsed_date=parse_date(raw_date)
                    if parsed_date is None:
                        out.issues.append(InputIssue("더존",ws.title,r,"기표일자",raw_date,"날짜를 해석할 수 없어 대상월 필터에서 제외")); continue
                    if not period.contains(parsed_date): continue

                formula_field=None
                for field,col in (("거래처코드",cols["vendor_code"]),("거래처명",cols["vendor_name"]),
                                  ("적요",cols["description"]),("차변",cols["debit"]),("대변",cols["credit"])):
                    cell=row[col-1]
                    if _is_formula(cell):
                        formula_field=(field,cell.value); break
                if formula_field:
                    out.issues.append(InputIssue("더존",ws.title,r,formula_field[0],formula_field[1],
                        "수식 셀은 계산값의 최신성을 보장할 수 없어 자동 처리하지 않음")); continue

                d=parse_amount(row[cols["debit"]-1].value); cr=parse_amount(row[cols["credit"]-1].value)
                suspicious=False
                for field,x in (("차변",d),("대변",cr)):
                    if x.quality==DataQuality.SUSPICIOUS:
                        out.issues.append(InputIssue("더존",ws.title,r,field,x.raw,x.reason)); suspicious=True
                if suspicious: continue

                debit=d.value or 0; credit=cr.value or 0
                if debit==0 and credit==0: continue
                if debit < 0 or credit < 0:
                    out.issues.append(InputIssue("더존",ws.title,r,"차변/대변",
                        f"{debit}/{credit}","음수 전표는 수정·역분개 가능성이 있어 자동 처리하지 않음")); continue
                if debit > 0 and credit > 0:
                    out.issues.append(InputIssue("더존",ws.title,r,"차변/대변",
                        f"{debit}/{credit}","한 행에 차변과 대변이 동시에 있어 자동 처리하지 않음")); continue

                code=_code_from_cell(row[cols["vendor_code"]-1])
                vendor_name=_text(row[cols["vendor_name"]-1].value)
                desc=_text(row[cols["description"]-1].value)
                row_invalid=False
                if not code:
                    out.issues.append(InputIssue("더존",ws.title,r,"거래처코드","", "거래처코드 없음")); row_invalid=True
                if not vendor_name:
                    out.issues.append(InputIssue("더존",ws.title,r,"거래처명","", "거래처명 없음")); row_invalid=True
                if not desc:
                    out.issues.append(InputIssue("더존",ws.title,r,"적요","", "적요 없음")); row_invalid=True
                if row_invalid: continue

                out.items.append(JournalLine(code,vendor_name,ac,an,desc,debit,credit,
                    _date(raw_date) if cols["date"] else "",r))
    finally:
        wb.close()
    if not out.items and not out.issues: raise ValueError("더존 파일에서 대상 미지급금 전표를 찾지 못했습니다.")
    return out


def read_prior(path: str|Path, owner: str = "", period: AccountingPeriod|None=None) -> ReadResult:
    wb=_load_workbook(path,read_only=True,data_only=False); out=ReadResult()
    verified_zero_groups=0
    groups=[("거래처코드","거래처 코드","코드"),("거래처명","거래처","업체명"),("적요","내역","내용"),("금액","미지급금","잔액")]
    try:
        # Generated month-end drafts containing unresolved reviews must not be
        # silently treated as approved opening balances in the next month.
        pending=next((ws for ws in wb.worksheets if ws.title=="검토필요"),None)
        if (pending is not None and pending.cell(1,10).value=="MONTH_END_DRAFT_REVIEW"
                and pending.max_row>1):
            raise ValueError(
                "자동 생성된 명세서에 검토필요 항목이 남아 있습니다. "
                "본문을 확인·수정한 뒤 검토필요 시트를 삭제하고 다음 달 대사를 진행해주세요."
            )
        selected=_select_prior_sheets(wb,groups,period)
        for ws,hr,vals in selected:
            out.recognized_sheets.append(ws.title)
            cols={
                "vendor_code":_first(vals,"거래처코드","거래처 코드","코드"),
                "vendor_name":_first(vals,"거래처명","거래처","업체명"),
                "date":_first(vals,"날짜","일자","기표일자"),
                "description":_first(vals,"적요","내역","내용"),
                "amount":_first(vals,"금액","미지급금","잔액")}
            required=("vendor_code","vendor_name","description","amount")
            missing=[x for x in required if not cols[x]]
            if missing:
                raise ValueError("명세서 필수 헤더를 안전하게 식별하지 못했습니다: "+", ".join(missing))
            out.detected_headers=[_text(v) for v in next(
                islice(ws.iter_rows(min_row=hr,values_only=True),1), ()
            )]

            # Some real statement sheets place vendor code/name only on the yellow subtotal row
            # after the detail rows. Buffer only contiguous detail rows and attach the vendor
            # identity when the following subtotal exactly equals their amount sum.
            pending=[]
            group_tainted=False
            flat_negative_codes=set()
            def flush_pending_as_issues(reason):
                nonlocal pending, group_tainted
                group_tainted=True
                for rec in pending:
                    out.issues.append(InputIssue("명세서",ws.title,rec["row"],"거래처",rec["description"],reason))
                pending=[]

            for r,row in enumerate(ws.iter_rows(min_row=hr+1,values_only=False),start=hr+1):
                code_cell=row[cols["vendor_code"]-1]
                name_cell=row[cols["vendor_name"]-1]
                desc_cell=row[cols["description"]-1]
                amount_cell=row[cols["amount"]-1]
                date_cell=row[cols["date"]-1] if cols["date"] else None

                code=_code_from_cell(code_cell)
                vendor_name=_text(name_cell.value)
                desc=_text(desc_cell.value)
                raw_date=date_cell.value if date_cell else None
                summary_label=next(
                    (v for v in (raw_date,desc,vendor_name) if _is_summary_label(v)),
                    None,
                )

                relevant_cells=[code_cell,name_cell,desc_cell,amount_cell]
                if date_cell is not None:
                    relevant_cells.append(date_cell)
                # A subtotal =SUM(E3:E15) is safe only when it sums exactly
                # this contiguous detail group. Compute it ourselves rather
                # than trusting a possibly stale formula cache.
                simple_sum=False
                if code and vendor_name and summary_label and pending and _is_formula(amount_cell):
                    column=get_column_letter(cols["amount"])
                    expected=f"=SUM({column}{pending[0]['row']}:{column}{pending[-1]['row']})"
                    actual=str(amount_cell.value).replace(" ", "").replace("$", "").upper()
                    simple_sum=(actual==expected)
                if any(_is_formula(cell) for cell in relevant_cells
                       if cell is not amount_cell or not simple_sum):
                    if pending:
                        flush_pending_as_issues("그룹 중간에 수식 셀이 있어 거래처를 안전하게 확정할 수 없음")
                    out.issues.append(InputIssue("명세서",ws.title,r,"수식",amount_cell.value,
                        "수식 셀은 계산값의 최신성을 보장할 수 없어 자동 대사하지 않음"))
                    if code and vendor_name and summary_label:
                        group_tainted=False
                    continue

                a=parse_amount(sum(rec["amount"] for rec in pending) if simple_sum else amount_cell.value)
                if not code and not vendor_name and not desc and (a.value or 0)==0:
                    if pending:
                        flush_pending_as_issues("거래처 소계 전에 빈 행이 있어 그룹 범위를 확정할 수 없음")
                    continue
                if a.quality==DataQuality.SUSPICIOUS:
                    if pending:
                        flush_pending_as_issues("그룹 중간에 확인이 필요한 금액 형식이 있어 거래처를 확정할 수 없음")
                    out.issues.append(InputIssue("명세서",ws.title,r,"금액",a.raw,a.reason))
                    if code and vendor_name and summary_label:
                        group_tainted=False
                    continue

                # Legacy presentation subtotal: vendor-name column itself contains only '소계'.
                if not code and not desc and _is_summary_label(vendor_name):
                    if pending:
                        flush_pending_as_issues("소계 행에 거래처코드/거래처명이 없어 상세 행의 거래처를 확정할 수 없음")
                    continue

                # Group subtotal row: vendor code/name + '소계' and exact sum of preceding details.
                if code and vendor_name and summary_label:
                    subtotal=a.value or 0
                    if group_tainted:
                        out.issues.append(InputIssue(
                            "명세서",ws.title,r,"소계",a.raw,
                            "거래처 상세 내역에 입력 오류가 있어 소계 확정 불가"
                        ))
                        pending=[]
                        group_tainted=False
                        continue
                    if not pending:
                        # A subtotal without detail rows is presentation-only.
                        continue
                    pending_total=sum(x["amount"] for x in pending)
                    if pending_total != subtotal:
                        reason=f"상세 합계 {pending_total:,.0f}원과 소계 {subtotal:,.0f}원이 달라 자동 대사하지 않음"
                        flush_pending_as_issues(reason)
                        out.issues.append(InputIssue("명세서",ws.title,r,"소계",a.raw,reason))
                        group_tainted=False
                        continue
                    if subtotal < 0:
                        flush_pending_as_issues("거래처 소계가 음수라 지급 초과/조정 여부 확인 필요")
                        out.issues.append(InputIssue("명세서",ws.title,r,"소계",a.raw,"거래처 순잔액 음수 확인 필요"))
                        group_tainted=False
                        continue

                    negatives=[rec for rec in pending if rec["amount"] < 0]
                    if negatives:
                        # Preserve a verified zero opening too: current-month debits
                        # may then create an overpayment that must not disappear.
                        if subtotal >= 0:
                            positive_total=sum(rec["amount"] for rec in pending if rec["amount"] > 0)
                            negative_total=sum(rec["amount"] for rec in negatives)
                            balance=PayableItem(
                                code,vendor_name,"전월 순미지급 잔액(음수 상세 포함)",subtotal,
                                pending[-1]["date"],r,
                                SourceRef(Path(path).name,ws.title,r,owner)
                            )
                            reason=(
                                f"양수 {positive_total:,.0f}원 + 음수 {negative_total:,.0f}원"
                                f" = 소계 {subtotal:,.0f}원 검증 완료. "
                                "음수 지급/조정의 개별 발생 건 상계가 불명확하여 자동 대사·이월 제외"
                            )
                            out.review_items.append((balance,reason))
                        # A zero closing balance has no outstanding payable to carry.
                        if subtotal == 0:
                            verified_zero_groups+=1
                    else:
                        for rec in pending:
                            out.items.append(PayableItem(
                                code,vendor_name,rec["description"],rec["amount"],
                                rec["date"],rec["row"],SourceRef(Path(path).name,ws.title,rec["row"],owner)
                            ))
                    pending=[]
                    group_tainted=False
                    continue

                # Flat format: every detail row already carries vendor identity.
                if code or vendor_name:
                    if pending:
                        flush_pending_as_issues("거래처 소계가 나오기 전에 다른 거래처 행이 시작되어 그룹을 확정할 수 없음")
                    row_invalid=False
                    if not code:
                        out.issues.append(InputIssue("명세서",ws.title,r,"거래처코드","", "거래처코드 없음")); row_invalid=True
                    if not vendor_name:
                        out.issues.append(InputIssue("명세서",ws.title,r,"거래처명","", "거래처명 없음")); row_invalid=True
                    if not desc:
                        out.issues.append(InputIssue("명세서",ws.title,r,"적요","", "적요 없음")); row_invalid=True
                    if (a.value or 0) <= 0:
                        if code and a.value is not None and a.value < 0:
                            flat_negative_codes.add(code)
                        out.issues.append(InputIssue("명세서",ws.title,r,"금액",a.raw,
                            "금액이 0 이하이며 소계 없는 음수·0원 행은 미지급 순잔액을 확정할 수 없어 자동 대사하지 않음")); row_invalid=True
                    if row_invalid:
                        continue
                    out.items.append(PayableItem(
                        code,vendor_name,desc,a.value,_date(raw_date) if date_cell else "",r,
                        SourceRef(Path(path).name,ws.title,r,owner)
                    ))
                    continue

                # Grouped format detail row: description+amount are present, vendor is supplied by
                # the following subtotal row. Keep it buffered; never guess the vendor.
                if desc and (a.value or 0)!=0:
                    pending.append({
                        "row":r,
                        "date":_date(raw_date) if date_cell else "",
                        "description":desc,
                        "amount":a.value,
                    })
                    continue

                if not desc and (a.value or 0)>0:
                    if pending:
                        flush_pending_as_issues("적요 없는 행이 그룹 중간에 있어 거래처를 확정할 수 없음")
                    out.issues.append(InputIssue("명세서",ws.title,r,"적요","", "적요 없음"))
                    continue

            if pending:
                flush_pending_as_issues("파일 끝까지 거래처 소계가 없어 상세 행의 거래처를 확정할 수 없음")
            # Without a verified subtotal, a negative flat row can affect
            # other rows for the same vendor; suppress their auto-matching.
            if flat_negative_codes:
                retained=[]
                for item in out.items:
                    if item.source.sheet==ws.title and item.vendor_code in flat_negative_codes:
                        out.issues.append(InputIssue(
                            "명세서",ws.title,item.source.row,"거래처",item.vendor_code,
                            "동일 거래처에 소계 없는 음수 내역이 있어 순잔액 확인 전 자동 대사 보류"
                        ))
                    else:
                        retained.append(item)
                out.items=retained
    finally:
        wb.close()
    if not out.items and not out.review_items and not out.issues and verified_zero_groups==0:
        target=f" {period.label}" if period else ""
        raise ValueError(f"명세서에서{target} 대사할 항목을 찾지 못했습니다.")
    return out

